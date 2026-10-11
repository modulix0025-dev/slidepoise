"""P-TEXTFIT gate: measure every edited text frame with the vendored ``mx_textfit`` before export.

Called from the hook at the top of OM ``app/editor/deck.py::apply()`` only when
MODULEX_MAG_TEXTFIT=1. Pure measurement: it never changes the deck, sizes or layout; it
raises :class:`TextDoesNotFit` (Arabic first, English second) when an edited frame overflows.

Unit of measurement: the whole text frame of a shape with every edited paragraph replaced.
With the full library (``paragraph_box`` present) the box, insets, theme fonts and the editor's
per-paragraph character cap come from the library; otherwise the frame is read with python-pptx.
Blocking verdicts: ``does_not_fit`` and ``missing_glyphs``. ``font_missing`` (the font is not
visible to OM's account) cannot be measured: it is returned as an unmeasured fact, not a refusal.
"""
from __future__ import annotations

import re
import sys
from collections.abc import Iterable
from pathlib import Path
from typing import Any

TEXT_RE = re.compile(r"^t-(\d{1,4})-(\d{1,7})-(\d{1,4})$")
MESSAGE_AR = "النص في صفحة {slide} أطول من المكان المخصّص له. من فضلك اختصره شوية وجرّب تاني."
MESSAGE_EN = "The text on page {slide} is longer than its box. Please shorten it and try again."
BLOCKING = {"does_not_fit", "missing_glyphs"}
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class TextDoesNotFit(ValueError):
    """Named refusal; ``.code``, ``.message_ar``, ``.message_en`` and ``.details`` for the editor."""

    code = "MAG_TEXT_OVERFLOW"

    def __init__(self, details: list[dict]):
        slides = ", ".join(str(d["slide"]) for d in details)
        lib_ar = "؛ ".join(d["ar"] for d in details if d.get("ar"))
        self.message_ar = MESSAGE_AR.format(slide=slides) + (f" ({lib_ar})" if lib_ar else "")
        self.message_en = MESSAGE_EN.format(slide=slides)
        super().__init__(self.message_ar + " / " + self.message_en)
        self.details = details


def load_lib():
    vendor = Path(__file__).resolve().parent / "vendor"
    if str(vendor) not in sys.path:
        sys.path.insert(0, str(vendor))
    import mx_textfit

    return mx_textfit


def _items(texts: Any) -> Iterable[tuple[str, str]]:
    if isinstance(texts, dict):
        return [(str(k), "" if v is None else str(v)) for k, v in texts.items()]
    out = []
    for row in texts or []:
        if isinstance(row, dict):
            out.append((str(row.get("id")), str(row.get("text", ""))))
        elif isinstance(row, (list, tuple)) and len(row) == 2:
            out.append((str(row[0]), str(row[1])))
    return out


def _blocking(verdict: dict) -> bool:
    if "verdict" in verdict:
        return verdict["verdict"] in BLOCKING
    return not verdict.get("fits", True)


def _walk(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _walk(shape.shapes)


def _fallback_frame(deck_in, slide_no: int, shape_id: int, first_para: int) -> dict | None:
    """Box, insets, paragraphs and font of one frame read with python-pptx (library without paragraph_box)."""
    from pptx import Presentation

    prs = Presentation(str(deck_in))
    if not 1 <= slide_no <= len(prs.slides):
        return None
    shape = next((s for s in _walk(prs.slides[slide_no - 1].shapes) if s.shape_id == shape_id and s.has_text_frame),
                 None)
    if shape is None:
        return None
    frame = shape.text_frame
    para = frame.paragraphs[min(first_para, len(frame.paragraphs) - 1)]
    size = family = None
    for run in para.runs:
        rpr = run._r.find(f"{_A}rPr")
        if rpr is None:
            continue
        if size is None and rpr.get("sz"):
            size = int(rpr.get("sz")) / 100
        for tag in ("cs", "latin"):
            el = rpr.find(f"{_A}{tag}")
            if family is None and el is not None and el.get("typeface"):
                family = el.get("typeface")
    insets = tuple(int(v) if v is not None else d for v, d in
                   ((frame.margin_left, 91440), (frame.margin_top, 45720), (frame.margin_right, 91440),
                    (frame.margin_bottom, 45720)))
    return {"box_emu": (int(shape.width or 0), int(shape.height or 0)), "paragraphs": [p.text for p in frame.paragraphs],
            "font": {"family": family, "size_pt": size}, "rules": {"inset_emu": insets}}


def check_texts(deck_in, texts, inv=None, *, lib: Any = None, rules: dict | None = None) -> list[dict]:
    """Return per-frame fit facts; raise :class:`TextDoesNotFit` if any edited frame overflows."""
    lib = lib or load_lib()
    edits: dict[tuple[int, int], dict[int, str]] = {}
    for tid, text in _items(texts):
        m = TEXT_RE.match(tid)
        if m:
            edits.setdefault((int(m.group(1)), int(m.group(2))), {})[int(m.group(3))] = text
    results, failures = [], []
    for (slide_no, shape_id), paras in sorted(edits.items()):
        first = min(paras)
        info = None
        if hasattr(lib, "paragraph_box"):
            try:
                pb = lib.paragraph_box(str(deck_in), f"t-{slide_no}-{shape_id}-{first}")
                info = {"box_emu": pb["box_emu"], "paragraphs": list(pb["shape_paragraphs"]), "font": pb["font"],
                        "rules": dict(pb.get("rules") or {})}
            except Exception:  # noqa: BLE001 - ParagraphNotFound etc.: the editor's own checks decide
                info = None
        if info is None:
            info = _fallback_frame(deck_in, slide_no, shape_id, first)
        if info is None:
            continue
        body = list(info["paragraphs"])
        originals = {i: (body[i] if 0 <= i < len(body) else "") for i in paras}
        for idx, new in paras.items():
            if 0 <= idx < len(body):
                body[idx] = new
        merged = {**info["rules"], **(rules or {})}
        merged.pop("max_chars", None)
        verdict = lib.fit("\n".join(body), info["box_emu"], info["font"], merged)
        if hasattr(lib, "char_limit") and verdict.get("verdict") not in ("missing_glyphs", "font_missing"):
            for idx, new in paras.items():
                limit = lib.char_limit(originals[idx])
                if len(new) > limit:
                    verdict = dict(verdict, fits=False, verdict="does_not_fit", reason="too_many_chars",
                                   ar=f"النص ده {len(new)} حرف والمسموح {limit} بس — قصّره شوية")
        row = {"slide": slide_no, "shape_id": shape_id, "paragraphs": sorted(paras),
               "verdict": verdict.get("verdict", "fits" if verdict.get("fits") else "does_not_fit"),
               "fits": bool(verdict.get("fits")), "ar": verdict.get("ar"), "en": verdict.get("en"),
               "method": verdict.get("method"), "unmeasured": verdict.get("verdict") == "font_missing"}
        results.append(row)
        if _blocking(verdict):
            failures.append(row)
    if failures:
        raise TextDoesNotFit(failures)
    return results
