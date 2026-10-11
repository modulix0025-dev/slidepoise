"""Build a one-slide copy of a deck and export it with OM's own ``export_pdf``; cache by content hash.

Cache key = sha256(deck bytes) + slide number → ``<cache>/<sha>-<slide>.pdf``. The deck on disk
is never modified (the one-slide copy is written to a temp file next to the cache).
A hidden slide is un-hidden in the copy so it can be previewed.
"""
from __future__ import annotations

import hashlib
import os
import tempfile
from collections.abc import Callable
from pathlib import Path


def single_slide_copy(src: Path, slide_no: int, dst: Path) -> None:
    from pptx import Presentation

    prs = Presentation(str(src))
    if not 1 <= slide_no <= len(prs.slides):
        raise ValueError(f"slide {slide_no} out of range 1..{len(prs.slides)}")
    sld_ids = prs.slides._sldIdLst
    keep = list(sld_ids)[slide_no - 1]
    for sld in list(sld_ids):
        if sld is not keep:
            prs.part.drop_rel(sld.rId)
            sld_ids.remove(sld)
    slide = prs.slides[0]
    if slide._element.get("show") == "0":
        del slide._element.attrib["show"]
    prs.save(str(dst))


def cached_preview(deck: Path, slide_no: int, cache_dir: Path, export_pdf: Callable) -> Path:
    deck = Path(deck)
    digest = hashlib.sha256(deck.read_bytes()).hexdigest()
    cache_dir.mkdir(parents=True, exist_ok=True)
    out = cache_dir / f"{digest}-{int(slide_no)}.pdf"
    if out.is_file() and out.stat().st_size > 0:
        return out
    fd, tmp_name = tempfile.mkstemp(suffix=".pptx", dir=str(cache_dir))
    os.close(fd)
    tmp = Path(tmp_name)
    tmp_pdf = out.with_name(out.name + ".tmp.pdf")
    try:
        single_slide_copy(deck, slide_no, tmp)
        export_pdf(str(tmp), str(tmp_pdf), 1)
        os.replace(tmp_pdf, out)
    finally:
        for p in (tmp, tmp_pdf):
            try:
                p.unlink()
            except OSError:
                pass
    return out
