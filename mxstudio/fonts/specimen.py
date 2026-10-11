"""Specimen book: one page per certified family + one page per type theme, as PPTX → PDF.

The deck is built with python-pptx (real text runs; Arabic paragraphs carry
``rtl="1"``, ``lang="ar-EG"`` and the ``a:cs`` typeface, like the fork's print
output) and exported through the operator's vendored ``pptx_to_pdf.convert``
(PowerPoint on the operator's PC, LibreOffice elsewhere). After export the PDF
is *measured*: which fonts are embedded (``FontFile*`` present, subset name)
and whether the extracted text equals the input
(:func:`mxstudio.validators.pdf_text.text_proof`). A LibreOffice export is
evidence for LibreOffice only — the PowerPoint export proof stays pending on
device. Typefaces are written by their **install name** (renamed ``MX Kids …``
builds where an RFN applies).
"""
from __future__ import annotations

import importlib.util
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Mm, Pt

from mxstudio.fonts.paths import ROOT

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PANGRAM_AR = "صِف خَلقَ خَودِ كَمِثلِ الشَّمسِ إِذ بَزَغَت يَحظى الضَّجيعُ بِها نَجلاءَ مِعطارِ"
PLAIN_AR = "ذهب أحمد وفاطمة إلى الحديقة ولعبا بالكرة حتى المغرب"
VOWELIZED_AR = "مُحَمَّدٌ وَفَاطِمَةُ يَلْعَبَانِ فِي الحَدِيقَةِ"
DIGITS_AR = "٠١٢٣٤٥٦٧٨٩ — 0123456789"
PANGRAM_EN = "The quick brown fox jumps over the lazy dog"
NAMES_AR = "عبد الرحمن · فاطمة الزهراء · نور الهدى · يوسف · مريم"
NAMES_EN = "Abdelrahman · Fatma Elzahraa · Nour Elhoda · Youssef · Mariam"
PAGE_MM = {"A4": (210, 297), "A5": (148, 210)}


@dataclass
class SpecimenFamily:
    family: str          # catalog family
    install_name: str    # what the deck references
    script: str          # "AR" | "EN"
    status: str = "certified_print"
    notes: list[str] = field(default_factory=list)


@dataclass
class SpecimenTheme:
    id: str
    name_ar: str
    name_en: str
    fonts: dict[str, str]  # slot → install name


def _set_run(run, text: str, typeface: str, size_pt: float, *, rtl: bool, bold: bool = False,
             colour: str = "1A1A1A") -> None:
    run.text = text
    font = run.font
    font.size = Pt(size_pt)
    font.bold = bold
    font.color.rgb = RGBColor.from_string(colour)
    rpr = run._r.get_or_add_rPr()  # noqa: SLF001
    rpr.set("lang", "ar-EG" if rtl else "en-US")
    for tag in ("latin", "ea", "cs"):
        node = rpr.find(f"{{{A}}}{tag}")
        if node is None:
            node = etree.SubElement(rpr, f"{{{A}}}{tag}")
        node.set("typeface", typeface)


def _line(slide, text: str, typeface: str, size_pt: float, top_mm: float, height_mm: float, width_mm: float, *,
          rtl: bool, bold: bool = False, colour: str = "1A1A1A") -> str:
    box = slide.shapes.add_textbox(Mm(12), Mm(top_mm), Mm(width_mm - 24), Mm(height_mm))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    ppr = paragraph._p.get_or_add_pPr()  # noqa: SLF001
    ppr.set("algn", "r" if rtl else "l")
    if rtl:
        ppr.set("rtl", "1")
    _set_run(paragraph.add_run(), text, typeface, size_pt, rtl=rtl, bold=bold, colour=colour)
    return text


def build_specimen(out_pptx: Path, families: list[SpecimenFamily], themes: list[SpecimenTheme], *,
                   page: str = "A4", label_font: str = "Tajawal") -> tuple[Path, list[list[str]]]:
    """Write the specimen deck; returns (path, the text strings placed on each page, in order)."""
    width_mm, height_mm = PAGE_MM[page]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(Mm(width_mm)), Emu(Mm(height_mm))
    blank = prs.slide_layouts[6]
    pages: list[list[str]] = []
    for fam in families:
        slide = prs.slides.add_slide(blank)
        texts: list[str] = []
        pages.append(texts)
        ar = fam.script == "AR"
        texts.append(_line(slide, fam.family, fam.install_name, 30, 10, 18, width_mm, rtl=False, bold=True))
        texts.append(_line(slide, f"install name: {fam.install_name} | {fam.status}", label_font, 9, 30, 8,
                           width_mm, rtl=False, colour="555555"))
        y = 42.0
        samples = ([(PANGRAM_AR, 22), (PLAIN_AR, 16), (VOWELIZED_AR, 18), (NAMES_AR, 16), (DIGITS_AR, 14)] if ar
                   else [(PANGRAM_EN, 22), (PANGRAM_EN.upper(), 16), (NAMES_EN, 16), ("0123456789 & ! ?", 14)])
        for text, size in samples:
            texts.append(_line(slide, text, fam.install_name, size, y, size * 0.9, width_mm, rtl=ar))
            y += size * 0.9 + 6
        for note in fam.notes[:4]:
            texts.append(_line(slide, note, label_font, 9, y, 7, width_mm, rtl=True, colour="8A1C1C"))
            y += 8
    for theme in themes:
        slide = prs.slides.add_slide(blank)
        texts = []
        pages.append(texts)
        f = theme.fonts
        texts.append(_line(slide, theme.name_ar, f["ar_title"], 34, 10, 20, width_mm, rtl=True, bold=True))
        texts.append(_line(slide, theme.name_en, f["en_title"], 26, 32, 16, width_mm, rtl=False, bold=True))
        texts.append(_line(slide, "لـيوسف", f["ar_dedication"], 22, 52, 14, width_mm, rtl=True))
        texts.append(_line(slide, "مريم", f["ar_name"], 30, 68, 18, width_mm, rtl=True, bold=True))
        texts.append(_line(slide, PLAIN_AR, f["ar_body"], 16, 90, 14, width_mm, rtl=True))
        texts.append(_line(slide, VOWELIZED_AR, f["ar_body_vowelized"], 16, 106, 14, width_mm, rtl=True))
        texts.append(_line(slide, PANGRAM_EN, f["en_body"], 14, 124, 12, width_mm, rtl=False))
    out_pptx = Path(out_pptx)
    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_pptx)
    return out_pptx, pages


@lru_cache(maxsize=1)
def _pptx_to_pdf():
    path = ROOT / "vendor" / "modulex" / "pptx_to_pdf.py"
    spec = importlib.util.spec_from_file_location("mx_vendor_pptx_to_pdf", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


def embedded_fonts(pdf: Path) -> list[dict]:
    """Every font in the PDF: base name (subset tag stripped), subset flag, and whether a font program is embedded."""
    import pikepdf
    out: dict[str, dict] = {}
    with pikepdf.open(str(pdf)) as document:
        for obj in document.objects:
            if not isinstance(obj, pikepdf.Dictionary) or obj.get("/Type") != pikepdf.Name("/FontDescriptor"):
                continue
            name = str(obj.get("/FontName", ""))[1:]
            base = re.sub(r"^[A-Z]{6}\+", "", name)
            embedded = any(key in obj for key in ("/FontFile", "/FontFile2", "/FontFile3"))
            item = out.setdefault(base, {"name": base, "subset": name != base, "embedded": False})
            item["embedded"] = item["embedded"] or embedded
    return sorted(out.values(), key=lambda item: item["name"])


def _squash(name: str) -> str:
    return re.sub(r"[^0-9a-z]", "", name.casefold())


def embedding_proof(pdf: Path, install_names: list[str]) -> dict:
    """Which install names have an embedded font program in the PDF (matched on the squashed PS-style name)."""
    fonts = embedded_fonts(pdf)
    embedded = {_squash(f["name"].split("-")[0]) for f in fonts if f["embedded"]} | {
        _squash(f["name"]) for f in fonts if f["embedded"]}
    found, missing = [], []
    for name in install_names:
        (found if any(e.startswith(_squash(name)) for e in embedded) else missing).append(name)
    return {"embedded": found, "missing": missing, "pdf_fonts": fonts}


def page_text_proofs(pdf: Path, pages: list[list[str]]) -> list[dict]:
    """text_proof per page (each page split into its own PDF), so a failure names the family/theme page."""
    import tempfile

    import pikepdf

    from mxstudio.validators.pdf_text import extract_text, text_proof
    proofs = []
    with pikepdf.open(str(pdf)) as document, tempfile.TemporaryDirectory(prefix="mx_specimen_") as tmp:
        if len(document.pages) != len(pages):
            return [{"page": 0, "ok": False, "detail": f"{len(document.pages)} PDF pages for {len(pages)} slides"}]
        for index, texts in enumerate(pages):
            single = pikepdf.new()
            single.pages.append(document.pages[index])
            path = Path(tmp) / f"p{index + 1}.pdf"
            single.save(str(path))
            proof = text_proof(texts, path)
            proof["multiset_equal"] = _multiset_equal(texts, extract_text(path))
            proofs.append({"page": index + 1, "title": texts[0] if texts else "", **proof})
    return proofs


_MIRROR = str.maketrans("()[]{}«»", "(([[{{««")


def _multiset_equal(expected: list[str], extracted: str) -> bool:
    """Relaxed proof: same multiset of letters/digits, marks and (mirror-folded) punctuation, ignoring order and
    whitespace. It absorbs extractor artefacts (spaces inserted inside vowelized words, ligature components in
    visual order, mirrored brackets in RTL) but still fails on any missing, extra or changed character."""
    from collections import Counter

    from mxstudio.validators.pdf_text import normalize_for_proof
    want_base, want_marks, want_punct = normalize_for_proof(" ".join(expected).translate(_MIRROR))
    got_base, got_marks, got_punct = normalize_for_proof(extracted.translate(_MIRROR))
    return Counter(want_base) == Counter(got_base) and want_marks == got_marks and want_punct == got_punct


def export_specimen(pptx: Path, pdf: Path, pages: list[list[str]], install_names: list[str], *,
                    engines: list[str] | None = None, timeout: int = 900) -> dict:
    """Export through the vendored converter, then prove embedding + text per page. Raises the converter's error."""
    record = _pptx_to_pdf().convert(pptx, pdf, engines=engines, timeout=timeout, expect_pages=len(pages))
    proofs = page_text_proofs(pdf, pages)
    return {"export": {k: record[k] for k in ("engine", "engine_label", "pages", "fonts", "has_text", "sha256")},
            "embedding": embedding_proof(pdf, install_names), "pages": proofs,
            "text_ok_pages": sum(1 for p in proofs if p["ok"]),
            "multiset_ok_pages": sum(1 for p in proofs if p.get("multiset_equal")),
            "multiset_failed_pages": [p["title"] for p in proofs if not p.get("multiset_equal")],
            "text_failed_pages": [{"page": p["page"], "title": p["title"], "missing_words": p.get("missing_words"),
                                   "extra_marks": p.get("extra_marks"), "missing_marks": p.get("missing_marks"),
                                   "missing_punct": p.get("missing_punct"), "extra_punct": p.get("extra_punct")}
                                  for p in proofs if not p["ok"]]}


def specimen_inputs(catalog, qa: dict, builds: dict, themes: dict) -> tuple[list[SpecimenFamily], list[SpecimenTheme]]:
    """From the catalog DB: certified families (install names from the office-build plan) and every type theme."""
    plans = {p["family"]: p for p in builds["families"]}
    families = []
    for entry in qa["families"]:
        if entry["status"] != "certified_print":
            continue
        plan = plans.get(entry["family"], {})
        notes = [r["ar"] for r in entry["reasons"] if r["severity"] == "limit"]
        families.append(SpecimenFamily(family=entry["family"], install_name=plan.get("install_family") or entry["family"],
                                       script=entry["script"], status=entry["status"], notes=notes))
    names = {p["family"]: p.get("install_family") or p["family"] for p in builds["families"]}
    slots = ("ar_title", "ar_body", "ar_body_vowelized", "ar_name", "ar_dedication", "en_title", "en_body")
    spec_themes = [SpecimenTheme(id=t["id"], name_ar=t["name_ar"], name_en=t["name_en"],
                                 fonts={s: names.get(t[s], t[s]) for s in slots}) for t in themes["themes"]]
    return families, spec_themes


def main(argv: list[str] | None = None) -> int:
    import argparse
    import json

    from mxstudio.fonts.catalog import default_catalog
    from mxstudio.fonts.paths import OFFICE_BUILDS_JSON, QA_JSON, THEMES_JSON

    parser = argparse.ArgumentParser(description="Generate the font specimen book (PPTX + PDF)")
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--engines", default="", help="e.g. soffice or powerpoint (default: operator order)")
    args = parser.parse_args(argv)
    load = lambda p: json.loads(p.read_text(encoding="utf-8"))  # noqa: E731
    families, themes = specimen_inputs(default_catalog(), load(QA_JSON), load(OFFICE_BUILDS_JSON), load(THEMES_JSON))
    pptx, pages = build_specimen(args.out_dir / "mx_font_specimen.pptx", families, themes)
    result = export_specimen(pptx, args.out_dir / "mx_font_specimen.pdf", pages,
                             sorted({f.install_name for f in families}),
                             engines=[e for e in args.engines.split(",") if e] or None)
    summary = {"pages_expected": len(pages), "families": len(families), "themes": len(themes),
               "export": result["export"], "embedded": len(result["embedding"]["embedded"]),
               "missing": result["embedding"]["missing"], "text_ok_pages": result["text_ok_pages"],
               "multiset_ok_pages": result["multiset_ok_pages"],
               "multiset_failed_pages": result["multiset_failed_pages"],
               "text_failed_pages": result["text_failed_pages"]}
    (args.out_dir / "specimen_proof.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1),
                                                      encoding="utf-8")
    print(json.dumps({k: summary[k] for k in ("pages_expected", "embedded", "text_ok_pages", "multiset_ok_pages",
                                              "multiset_failed_pages")}, ensure_ascii=False),
          "missing:", summary["missing"][:20])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
