"""Place shapes into PowerPoint as native, editable objects (python-pptx + DrawingML).

Rule: every shape lands as ``p:sp`` with ``a:custGeom`` or ``a:prstGeom`` (+ ``a:avLst``), and
text warps as ``a:bodyPr/a:prstTxWarp`` - never as pictures. Multi-layer generator output
becomes a ``p:grpSp`` of editable shapes. ``scene_freeform`` emits the authored-path object the
fork's renderer (``slidepoise/runtime/js/scene_to_pptx.mjs`` ``addFreeform``) already accepts.
"""
from __future__ import annotations

import zipfile
from pathlib import Path as FsPath

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from . import presets, warps
from .custgeom import A_NS, CustomGeometry, validate_custgeom_xml
from .generators import GeneratedShape

EMU_PER_MM = 36000
_A = "{%s}" % A_NS
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def mm(v: float) -> int:
    return int(round(v * EMU_PER_MM))


def new_presentation(width_mm: float = 148.0, height_mm: float = 210.0) -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(mm(width_mm))
    prs.slide_height = Emu(mm(height_mm))
    return prs


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rgb(color: str) -> RGBColor:
    c = color.lstrip("#")
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    return RGBColor.from_string(c.upper())


def _style(shape, fill: str | None, line: str | None, line_w_pt: float) -> None:
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(line_w_pt)
    else:
        shape.line.fill.background()


def _swap_geometry(shape, new_geom_xml: str) -> None:
    """Replace the placeholder geometry and drop the theme style reference (no stray shadows)."""
    sp_pr = shape._element.spPr
    old = sp_pr.find(_A + "prstGeom")
    if old is None:
        old = sp_pr.find(_A + "custGeom")
    new = etree.fromstring(new_geom_xml)
    old.addprevious(new)
    sp_pr.remove(old)
    style = shape._element.find(_P + "style")
    if style is not None:
        shape._element.remove(style)


def add_custgeom(shapes, geometry: CustomGeometry, x: int, y: int, w: int, h: int, *, fill: str | None = "#000000",
                 line: str | None = None, line_w_pt: float = 1.0, name: str | None = None):
    """Add an editable freeform; ``shapes`` is a slide's or a group's shape collection."""
    sp = shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    _swap_geometry(sp, geometry.to_xml(with_ns=True))
    _style(sp, fill, line, line_w_pt)
    if name:
        sp.name = name
    return sp


def add_preset(shapes, prst: str, x: int, y: int, w: int, h: int, *, adj: dict | None = None,
               fill: str | None = "#4F81BD", line: str | None = "#1F3864", line_w_pt: float = 1.0,
               text: str | None = None, name: str | None = None):
    sp = shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    _swap_geometry(sp, presets.prst_geom_xml(prst, adj, with_ns=True))
    _style(sp, fill, line, line_w_pt)
    if text:
        sp.text_frame.text = text
    sp.name = name or prst
    return sp


def add_warped_text(shapes, text: str, warp: str, x: int, y: int, w: int, h: int, *, adj: dict | None = None,
                    font: str = "Arial", size_pt: float = 36, color: str = "#C0392B", rtl: bool = False):
    tb = shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(size_pt)
    run.font.name = font
    run.font.color.rgb = _rgb(color)
    if rtl:
        tf.paragraphs[0]._p.get_or_add_pPr().set("rtl", "1")
    body_pr = tf._txBody.find(_A + "bodyPr")
    warp_el = etree.fromstring(warps.prst_tx_warp_xml(warp, adj, with_ns=True))
    body_pr.insert(0, warp_el)
    tb.name = f"warp:{warp}"
    return tb


DEFAULT_PALETTE = {"primary": "#E63946", "secondary": "#457B9D", "accent": "#F4A261", "cut": "#FFFFFF",
                   "art": "#222222"}
_CYCLE = ["#E63946", "#F4A261", "#E9C46A", "#2A9D8F", "#457B9D", "#9B5DE5", "#EF476F", "#06D6A0"]


def role_colour(role: str, palette: dict | None = None) -> str:
    pal = {**DEFAULT_PALETTE, **(palette or {})}
    if role in pal:
        return pal[role]
    if "_" in role and role.rsplit("_", 1)[1].isdigit():
        return _CYCLE[int(role.rsplit("_", 1)[1]) % len(_CYCLE)]
    return pal["primary"]


def add_generated(slide, gen: GeneratedShape, x: int, y: int, w: int, h: int, *, palette: dict | None = None,
                  line: str | None = None, line_w_pt: float = 1.0):
    """One shape per layer; several layers become one group (still individually editable)."""
    if len(gen.layers) == 1:
        layer = gen.layers[0]
        stroke_only = all(p.fill == "none" for p in layer.geometry.parts)
        has_strokes = any(p.stroke for p in layer.geometry.parts)
        return add_custgeom(slide.shapes, layer.geometry, x, y, w, h,
                            fill=None if stroke_only else role_colour(layer.role, palette),
                            line=line or ("#222222" if any(p.fill == "none" for p in layer.geometry.parts)
                                          and has_strokes else None),
                            line_w_pt=line_w_pt, name=f"{gen.name}:{layer.role}")
    grp = slide.shapes.add_group_shape()
    for layer in gen.layers:
        add_custgeom(grp.shapes, layer.geometry, x, y, w, h, fill=role_colour(layer.role, palette), line=line,
                     line_w_pt=line_w_pt, name=f"{gen.name}:{layer.role}")
    grp.name = gen.name
    return grp


def scene_freeform(geometry: CustomGeometry, bbox_px: list[float], object_id: str, style: dict | None = None) -> dict:
    """Scene object for ``scene_to_pptx.mjs`` (``kind: freeform`` + ``path_commands_px``; M/L/C/Z only)."""
    x, y, w, h = [float(v) for v in bbox_px]
    sx, sy = w / geometry.width, h / geometry.height
    cmds = []
    for part in geometry.parts:
        for c in part.path.only_mlcz().commands:
            op = c[0]
            if op == "Z":
                cmds.append({"op": "Z"})
            elif op in ("M", "L"):
                cmds.append({"op": op, "point": [x + c[1] * sx, y + c[2] * sy]})
            elif op == "C":
                cmds.append({"op": "C", "control1": [x + c[1] * sx, y + c[2] * sy],
                             "control2": [x + c[3] * sx, y + c[4] * sy], "point": [x + c[5] * sx, y + c[6] * sy]})
    return {"id": object_id, "kind": "freeform", "bbox_px": [x, y, w, h], "path_commands_px": cmds,
            "style": style or {"fill": "#E63946", "stroke": "none"}}


def validate_pptx(path: str | FsPath) -> dict:
    """Facts about every shape geometry in a deck + named findings (empty findings = valid)."""
    findings: list[str] = []
    counts = {"custGeom": 0, "prstGeom": 0, "prstTxWarp": 0, "pic": 0}
    prst_ok = set(presets.ECMA_PRESET_SHAPES)
    warp_ok = set(warps.ECMA_TEXT_WARPS)
    with zipfile.ZipFile(path) as z:
        for name in sorted(n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")):
            root = etree.fromstring(z.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
            for cg in root.iter(_A + "custGeom"):
                counts["custGeom"] += 1
                for f in validate_custgeom_xml(etree.tostring(cg)):
                    findings.append(f"{name}: {f}")
            for pg in root.iter(_A + "prstGeom"):
                counts["prstGeom"] += 1
                if pg.get("prst") not in prst_ok:
                    findings.append(f"{name}: unknown prst {pg.get('prst')}")
            for tw in root.iter(_A + "prstTxWarp"):
                counts["prstTxWarp"] += 1
                if tw.get("prst") not in warp_ok:
                    findings.append(f"{name}: unknown prstTxWarp {tw.get('prst')}")
                body = tw.getparent()
                if etree.QName(body).localname != "bodyPr" or body.index(tw) != 0:
                    findings.append(f"{name}: prstTxWarp must be the first child of a:bodyPr")
            counts["pic"] += len(root.findall(".//" + _P + "pic"))
    return {"counts": counts, "findings": findings}


def build_proof_deck(path: str | FsPath, *, svg_files: list | None = None, seed: int = 1) -> dict:
    """Deck with every preset (default handles), every text warp, every generator and optional
    converted SVGs - the artefact the Proof Runner opens in PowerPoint on the operator's PC."""
    from .generators import REGISTRY, generate
    from .svg_convert import svg_to_shapes

    prs = new_presentation(210, 148)
    W, H = prs.slide_width, prs.slide_height
    cols, rows = 8, 5
    cw, ch = W // cols, (H - mm(14)) // rows
    allp = presets.all_presets()
    per = cols * rows
    for start in range(0, len(allp), per):
        slide = blank_slide(prs)
        tb = slide.shapes.add_textbox(mm(4), mm(2), W - mm(8), mm(10))
        tb.text_frame.text = f"Preset geometries {start + 1}-{min(start + per, len(allp))} / {len(allp)}"
        for i, p in enumerate(allp[start:start + per]):
            x = (i % cols) * cw + mm(3)
            y = mm(14) + (i // cols) * ch + mm(2)
            add_preset(slide.shapes, p.prst, x, y, cw - mm(6), ch - mm(8), name=p.prst)
            lab = slide.shapes.add_textbox(x, y + ch - mm(7), cw - mm(6), mm(5))
            lab.text_frame.text = p.prst
            lab.text_frame.paragraphs[0].runs[0].font.size = Pt(6)
    allw = warps.all_warps()
    per_w = 12
    for start in range(0, len(allw), per_w):
        slide = blank_slide(prs)
        for i, wp in enumerate(allw[start:start + per_w]):
            x = (i % 4) * (W // 4) + mm(3)
            y = (i // 4) * (H // 3) + mm(3)
            add_warped_text(slide.shapes, "Hello مرحبا", wp.prst, x, y, W // 4 - mm(6), H // 3 - mm(12), size_pt=20)
            lab = slide.shapes.add_textbox(x, y + H // 3 - mm(10), W // 4 - mm(6), mm(5))
            lab.text_frame.text = wp.prst
            lab.text_frame.paragraphs[0].runs[0].font.size = Pt(7)
    gens = list(REGISTRY)
    per_g = 6
    for start in range(0, len(gens), per_g):
        slide = blank_slide(prs)
        tb = slide.shapes.add_textbox(mm(4), mm(2), W - mm(8), mm(10))
        tb.text_frame.text = "Procedural generators (seed %d)" % seed
        for i, name in enumerate(gens[start:start + per_g]):
            params = {}
            if name == "sticker_outline":
                params = {"source": generate("blob", seed).layers[0].geometry.to_json(), "include_art": True}
            g = generate(name, seed, params)
            bw, bh = W // 3 - mm(8), (H - mm(14)) // 2 - mm(10)
            scale = min(bw / g.width, bh / g.height)
            x = (i % 3) * (W // 3) + mm(4)
            y = mm(14) + (i // 3) * ((H - mm(14)) // 2)
            add_generated(slide, g, x, y, int(g.width * scale), int(g.height * scale))
            lab = slide.shapes.add_textbox(x, y + bh + mm(1), bw, mm(5))
            lab.text_frame.text = name
            lab.text_frame.paragraphs[0].runs[0].font.size = Pt(7)
    for svg_path in svg_files or []:
        res = svg_to_shapes(FsPath(svg_path).read_text(encoding="utf-8"))
        slide = blank_slide(prs)
        tb = slide.shapes.add_textbox(mm(4), mm(2), W - mm(8), mm(10))
        tb.text_frame.text = f"SVG -> custGeom: {FsPath(svg_path).name}"
        scale = min((W - mm(20)) / res.width, (H - mm(24)) / res.height)
        for s in res.shapes:
            local, (bx, by, bw, bh) = s.geometry.cropped()
            add_custgeom(slide.shapes, local, mm(10) + int(bx * scale), mm(16) + int(by * scale),
                         max(1, int(bw * scale)), max(1, int(bh * scale)), fill=s.fill or None,
                         line=s.stroke if s.stroke and s.stroke_width else None,
                         line_w_pt=max(0.25, s.stroke_width * scale / 12700))
    prs.save(str(path))
    return {"slides": len(prs.slides), **validate_pptx(path)}
