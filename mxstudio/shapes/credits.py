"""Credits page and licence report for a template/book.

Rules
* Every asset is listed in the licence report with its full provenance.
* Assets whose licence requires attribution (CC-BY-4.0, or a source that says
  ``AttributionRequired``) get a credits-page line; the report fails with a named finding if
  any credited asset lacks author, licence URL or source URL.
* Lines are Egyptian Arabic first, English second. Provenance text is printed as data.
"""
from __future__ import annotations

from .discovery.base import ORIGINAL_WORK, Provenance, clean_text, licence_gate

LICENCE_LABEL = {"CC-BY-4.0": "CC BY 4.0", "CC0-1.0": "CC0 1.0", "PDM": "Public Domain Mark"}


def credit_line(p: Provenance) -> dict:
    lic = LICENCE_LABEL.get(p.licence, p.licence)
    title = clean_text(p.title) or clean_text(p.source_id) or "—"
    author = clean_text(p.author) or "—"
    ar = f"«{title}» — تصميم: {author} — الرخصة: {lic} ({p.licence_url}) — المصدر: {p.source_url}"
    en = f"“{title}” by {author} — licence: {lic} ({p.licence_url}) — source: {p.source_url}"
    return {"ar": ar, "en": en, "licence": p.licence, "source_url": p.source_url, "sha256": p.sha256}


def licence_report(provenances: list[Provenance]) -> dict:
    """All assets + the credits-page entries + named findings."""
    rows, credits, findings = [], [], []
    for p in provenances:
        decision = licence_gate(p.licence, attribution_required=p.credits_required)
        rows.append({**p.to_json(), "decision": decision.status})
        missing = p.missing()
        if missing:
            findings.append({"code": "provenance_incomplete", "asset": p.sha256[:12], "missing": missing})
        if not decision.printable:
            findings.append({"code": "licence_not_printable", "asset": p.sha256[:12], "licence": p.licence})
        if decision.credits_required:
            need = [k for k in ("author", "licence_url", "source_url") if not str(getattr(p, k) or "").strip()]
            if need:
                findings.append({"code": "credit_incomplete", "asset": p.sha256[:12], "missing": need})
            credits.append(credit_line(p))
    credits.sort(key=lambda c: (c["licence"], c["en"]))
    originals = sum(1 for p in provenances if p.licence == ORIGINAL_WORK)
    return {"assets": rows, "credits": credits, "credits_page_required": bool(credits), "findings": findings,
            "summary": {"total": len(provenances), "originals": originals, "credited": len(credits)}}


def report_markdown(report: dict) -> str:
    lines = ["# تقرير التراخيص — Licence report", "",
             f"إجمالي الأصول: {report['summary']['total']} — Total assets: {report['summary']['total']}", ""]
    lines += ["| المصدر Source | العنوان Title | المؤلف Author | الرخصة Licence | sha256 |", "|---|---|---|---|---|"]
    for r in report["assets"]:
        lines.append(f"| {clean_text(r['source'])} | {clean_text(r.get('title'))} | {clean_text(r['author'])} | "
                     f"{r['licence']} | {r['sha256'][:16]} |")
    if report["credits"]:
        lines += ["", "## صفحة الاعتمادات — Credits", ""]
        for c in report["credits"]:
            lines += [f"- {c['ar']}", f"  {c['en']}"]
    return "\n".join(lines) + "\n"


def append_credits_slide(prs, report: dict, *, font: str = "Arial", size_pt: float = 9.0):
    """Append a credits slide (Arabic first) when the report requires one; returns it or None."""
    from pptx.util import Emu, Pt

    if not report["credits"]:
        return None
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    margin = int(prs.slide_width * 0.06)
    tb = slide.shapes.add_textbox(Emu(margin), Emu(margin), Emu(prs.slide_width - 2 * margin),
                                  Emu(prs.slide_height - 2 * margin))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = "الاعتمادات — Credits"
    tf.paragraphs[0].runs[0].font.size = Pt(size_pt * 1.6)
    for c in report["credits"]:
        for text in (c["ar"], c["en"]):
            para = tf.add_paragraph()
            para.text = text
            para.runs[0].font.size = Pt(size_pt)
            para.runs[0].font.name = font
    tb.name = "credits"
    return slide
