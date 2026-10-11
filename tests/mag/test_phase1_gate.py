"""Phase 1 gate: synthetic pages built to Appendix A anatomy, judged by the operator's own rules.

* Console rule  (`deck_picture_slots`, verbatim) returns exactly the declared slot map;
* OM rule       (`replace_in_shape`, verbatim) replaces exactly the expected paragraphs;
* every hero name lies whole inside one run;
* the page is the physical A5/A4 canvas incl. bleed, paragraphs are RTL with one pPr;
* python-pptx sees picture placeholders and injection keeps the slot map;
* LibreOffice exports a text-bearing PDF (vendored `pptx_to_pdf`) whose words and
  harakat equal the deck's;
* negative controls prove each check can fail.
"""
from __future__ import annotations

import os
import re
import sys
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from mxstudio.validators.operator_ports import deck_picture_slots, name_run_integrity, replay_generate_story
from mxstudio.validators.pdf_text import text_proof
from mxstudio.validators.synthetic_story import build_spec

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "vendor" / "modulex"))
import pptx_to_pdf  # noqa: E402  (vendored, byte-identical)

CASES = [("ar_boy", "A5"), ("ar_girl", "A5"), ("ar_boy", "A4")]


def _xml(deck: Path, part: str) -> str:
    with zipfile.ZipFile(deck) as archive:
        return archive.read(part).decode("utf-8")


def _deck_texts(deck: Path) -> list[str]:
    return [p.text for slide in Presentation(deck).slides for shape in slide.shapes if shape.has_text_frame
            for p in shape.text_frame.paragraphs]


@pytest.mark.parametrize("variant,size", CASES)
def test_console_slot_rule_matches_declared_map(render_deck, variant, size):
    deck, spec = render_deck(variant, size)
    assert deck_picture_slots(deck) == spec["_expected"]["picture_slots"]


@pytest.mark.parametrize("variant,size", CASES)
def test_om_replacement_rule_counts_exactly(render_deck, variant, size):
    deck, spec = render_deck(variant, size)
    base = spec["_expected"]["base_name"]
    result = replay_generate_story(deck, base, "عبد الرحمن")
    assert result["step1"] == spec["_expected"]["step1_replacements"]
    texts = [t for t in result["paragraph_texts"] if t]
    assert all(base not in t for t in texts)
    assert texts[-1] == "بكل الحب تهدى لـعبد الرحمن"
    assert all(item["ok"] for item in name_run_integrity(deck, base))


@pytest.mark.parametrize("variant,size", CASES)
def test_physical_page_rtl_and_structure(render_deck, variant, size):
    deck, spec = render_deck(variant, size)
    canvas = spec["page"]["canvas_mm"]
    sld = re.search(r'<p:sldSz cx="(\d+)" cy="(\d+)"', _xml(deck, "ppt/presentation.xml"))
    assert (int(sld.group(1)), int(sld.group(2))) == (round(canvas[0] * 36000), round(canvas[1] * 36000))
    for number in (1, 2, 3):
        slide = _xml(deck, f"ppt/slides/slide{number}.xml")
        for paragraph in re.findall(r"<a:p>.*?</a:p>", slide, re.S):
            ppr = re.findall(r"<a:pPr[^>]*>", paragraph)
            assert len(ppr) == 1 and 'rtl="1"' in ppr[0] and paragraph.index("<a:pPr") == len("<a:p>")
            assert 'lang="en-US"' not in paragraph and 'lang="ar-EG"' in paragraph
            assert "<a:cs typeface=" in paragraph


def test_full_bleed_placeholder_reaches_every_edge(render_deck):
    deck, spec = render_deck("ar_boy", "A5")
    slide = _xml(deck, "ppt/slides/slide4.xml")
    offsets = [tuple(map(int, m)) for m in re.findall(r'<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"', slide)
               if int(m[2]) > 0]  # skip the group frame of the shape tree
    assert len(offsets) == 2
    width, height = (round(v * 36000) for v in spec["page"]["canvas_mm"])
    xs = [(x, x + cx) for x, _, cx, _ in offsets]
    assert all(left == 0 and abs(right - width) <= 400 for left, right in xs)  # < 0.012 mm
    assert max(y + cy for _, y, _, cy in offsets) >= height - 400


def test_python_pptx_injection_keeps_slots(render_deck, tmp_path):
    deck, spec = render_deck("ar_boy", "A5")
    image = tmp_path / "scene.png"
    Image.new("RGB", (900, 800), (40, 120, 200)).save(image)
    prs = Presentation(deck)
    kinds = []
    for slide in prs.slides:
        for placeholder in slide.placeholders:
            kinds.append(type(placeholder).__name__)
            placeholder.insert_picture(str(image))
    out = tmp_path / "injected.pptx"
    prs.save(out)
    assert kinds == ["PicturePlaceholder"] * 4
    assert deck_picture_slots(out) == spec["_expected"]["picture_slots"]
    assert _xml(out, "ppt/slides/slide1.xml").count("<p:pic>") == 1


@pytest.mark.parametrize("variant,size", CASES)
def test_libreoffice_pdf_carries_every_word_and_haraka(render_deck, soffice, fixture_fontconfig, tmp_path,
                                                       monkeypatch, variant, size):
    deck, spec = render_deck(variant, size)
    monkeypatch.setenv("FONTCONFIG_FILE", str(fixture_fontconfig))
    monkeypatch.setenv("MODULEX_SOFFICE", soffice)
    pdf = tmp_path / f"{variant}_{size}.pdf"
    record = pptx_to_pdf.convert(deck, pdf, expect_pages=spec["_expected"]["slides"], engines=["soffice"])
    assert record["has_text"] and record["pages"] == spec["_expected"]["slides"]
    assert "page_mismatch" not in record
    proof = text_proof(_deck_texts(deck), pdf)
    assert proof["ok"], proof


def test_negative_control_text_proof_detects_a_missing_word(render_deck, soffice, fixture_fontconfig, tmp_path,
                                                           monkeypatch):
    deck, _ = render_deck("ar_boy", "A5")
    monkeypatch.setenv("FONTCONFIG_FILE", str(fixture_fontconfig))
    pdf = tmp_path / "proof.pdf"
    pptx_to_pdf.convert(deck, pdf, engines=["soffice"])
    texts = _deck_texts(deck)
    assert not text_proof(texts + ["سمكة"], pdf)["ok"]
    assert not text_proof([t.replace("كثيرًا", "كثيرا") for t in texts], pdf)["marks_equal"]


def test_negative_control_slot_rule_sees_a_missing_placeholder(render_deck):
    spec = build_spec("ar_boy", "A5")
    spec["slides"][3]["objects"].pop()  # drop one back-page placeholder
    deck, _ = render_deck(spec=spec, name="missing_placeholder")
    assert deck_picture_slots(deck) == {1: 1, 2: 1, 4: 1} != spec["_expected"]["picture_slots"]


def test_libreoffice_resave_destroys_placeholders(render_deck, soffice, tmp_path):
    """Measured rule: LibreOffice may EXPORT the deck, but must never re-save a template."""
    deck, _ = render_deck("ar_boy", "A5")
    import subprocess
    profile = (tmp_path / "profile").as_uri()
    subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pptx",
                    "--outdir", str(tmp_path / "rt"), str(deck)], capture_output=True, timeout=300, check=False,
                   env={**os.environ})
    resaved = tmp_path / "rt" / deck.name
    assert resaved.is_file()
    assert deck_picture_slots(resaved) == {}
