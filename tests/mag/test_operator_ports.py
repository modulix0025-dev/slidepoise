"""The operator's own rules, run verbatim, accept what the Studio builds — and reject what they should."""
from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm

from mxstudio.validators import operator_ports as ports

PORTS = Path(ports.__file__)
PINNED = {
    "OM app/services/story_engine.py (Appendix B.1, verbatim)":
        "56b3030b09002b0e39791b85d4174138bb363b4a8cd6c4ac7fdc877ec24c7a58",
    "Console magazine_pptx.py (Appendix B.4, verbatim)":
        "234be0932f53854ae1b2fc05b6086092dbcd45324d598de343aa99771158bb7a",
}


def _blocks() -> dict[str, str]:
    text = PORTS.read_text(encoding="utf-8")
    return {m.group(1): m.group(2) for m in
            re.finditer(r"# ---- PORT BEGIN: ([^\n]*?) ----\n(.*?)# ---- PORT END", text, re.S)}


def test_verbatim_blocks_are_pinned():
    blocks = _blocks()
    assert set(blocks) == set(PINNED)
    for name, body in blocks.items():
        assert hashlib.sha256(body.encode()).hexdigest() == PINNED[name], name


def _deck(tmp_path, paragraphs: list[list[str]], name="t.pptx") -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(5))
    frame = box.text_frame
    for index, runs in enumerate(paragraphs):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        for text in runs:
            para.add_run().text = text
    path = tmp_path / name
    prs.save(path)
    return path


def test_replay_counts_paragraphs_like_om(tmp_path):
    deck = _deck(tmp_path, [["ذهب إياد إلى البحر، وكان إياد سعيدًا"], ["لا اسم هنا"], ["بكل الحب تهدى لـإياد"]])
    result = ports.replay_generate_story(deck, "إياد", "محمد")
    assert result["step1"] == 2  # paragraphs, not occurrences
    assert "ذهب محمد إلى البحر، وكان محمد سعيدًا" in result["paragraph_texts"]
    assert "بكل الحب تهدى لـمحمد" in result["paragraph_texts"]


def test_step2_key_is_consumed_by_step1_when_the_name_is_literal(tmp_path):
    """Why the factory writes «لـ<name>»: an entry keyed on «لإياد» never fires after STEP 1."""
    deck = _deck(tmp_path, [["بكل الحب تهدى لإياد"]])
    result = ports.replay_generate_story(deck, "إياد", "محمد", {"بكل الحب تهدى لإياد": "بكل الحب تهدى لـ{child_name}"})
    assert result["step1"] == 1
    assert result["step2"] == {"بكل الحب تهدى لإياد": 0}
    assert result["paragraph_texts"] == ["بكل الحب تهدى لمحمد"]


def test_step2_fires_when_the_template_stores_a_ligature_form(tmp_path):
    """A lam-alef presentation form (U+FEF9 «ﻹ») hides the name from STEP 1; STEP 2 must catch it."""
    stored = "بكل الحب تهدى ﻹياد"
    deck = _deck(tmp_path, [[stored]])
    result = ports.replay_generate_story(deck, "إياد", "محمد", {stored: "بكل الحب تهدى لـ{child_name}"})
    assert result["step1"] == 0
    assert result["step2"] == {stored: 1}
    assert result["paragraph_texts"] == ["بكل الحب تهدى لـمحمد"]


def test_name_split_across_runs_is_flagged_and_loses_formatting(tmp_path):
    """Negative control: OM's cross-run fallback still replaces, but collapses runs."""
    deck = _deck(tmp_path, [["كان ", "إي", "اد", " سعيدًا"]])
    integrity = ports.name_run_integrity(deck, "إياد")
    assert integrity == [{"slide": 1, "shape": integrity[0]["shape"], "paragraph": 0, "occurrences": 1,
                          "whole_run_occurrences": 0, "ok": False}]
    out = tmp_path / "out.pptx"
    ports.replay_generate_story(deck, "إياد", "محمد", out_path=out)
    runs = [run.text for run in Presentation(out).slides[0].shapes[0].text_frame.paragraphs[0].runs]
    assert runs == ["كان محمد سعيدًا", "", "", ""]


def test_deck_picture_slots_counts_empty_and_filled(tmp_path):
    slide_xml = (b'<p:sld xmlns:p="p"><p:sp><p:nvSpPr><p:nvPr><p:ph type="pic" idx="1"/></p:nvPr></p:nvSpPr></p:sp>'
                 b'<p:pic><p:nvPicPr><p:nvPr><p:ph type="pic" idx="2"/></p:nvPr></p:nvPicPr></p:pic>'
                 b'<p:sp><p:nvSpPr><p:nvPr><p:ph type="body" idx="3"/></p:nvPr></p:nvSpPr></p:sp></p:sld>')
    path = tmp_path / "slots.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ppt/slides/slide3.xml", slide_xml)
        archive.writestr("ppt/slides/slide1.xml", b"<p:sld/>")
    assert ports.deck_picture_slots(path) == {3: 2}
    assert ports.deck_picture_slots(tmp_path / "missing.pptx") == {}
