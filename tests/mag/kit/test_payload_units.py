"""Unit facts of the operator-side payloads (no hook involved) + compatibility with the real mxlibs."""
from __future__ import annotations

import importlib
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

from mxstudio.kit.patches.base import PatchError, lib_files
from mxstudio.kit.patches.payloads.preflight2.hook import append_checks, profile_from, wrap_run
from mxstudio.kit.patches.payloads.preview.preview import cached_preview, single_slide_copy
from mxstudio.kit.patches.payloads.proof.approval import record, status
from mxstudio.kit.patches.payloads.safezone.geometry import geometry
from mxstudio.kit.patches.payloads.textfit.gate import TextDoesNotFit, check_texts
from mxstudio.validators.operator_ports import deck_picture_slots

REPO = Path(__file__).resolve().parents[3]


def _deck(facts):
    return facts["om"] / "orders" / "ORD-REPLICA-1" / "deck.pptx"


def _standin_textfit(standin_libs):
    spec = importlib.util.spec_from_file_location("standin_textfit", standin_libs / "mx_textfit" / "__init__.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _first_text_id(deck):
    from pptx import Presentation

    prs = Presentation(str(deck))
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return f"t-1-{shape.shape_id}-0"
    raise AssertionError("no text")


def test_textfit_gate_with_standin_lib(replica, standin_libs):
    facts = replica()
    deck = _deck(facts)
    lib = _standin_textfit(standin_libs)
    tid = _first_text_id(deck)
    rows = check_texts(deck, {tid: "نص قصير"}, None, lib=lib)
    assert rows[0]["fits"] and rows[0]["slide"] == 1
    with pytest.raises(TextDoesNotFit) as err:
        check_texts(deck, [{"id": tid, "text": "كلمة " * 200}], None, lib=lib)
    assert err.value.code == "MAG_TEXT_OVERFLOW" and err.value.message_ar.startswith("النص في صفحة 1")
    assert check_texts(deck, {"i-1-2": "x", "garbage": "y"}, None, lib=lib) == []


def test_textfit_gate_with_the_real_mx_textfit(replica):
    real = pytest.importorskip("mx_textfit")
    for name in ("fit", "paragraph_box", "char_limit"):
        assert hasattr(real, name), name
    facts = replica()
    deck = _deck(facts)
    tid = _first_text_id(deck)
    try:
        rows = check_texts(deck, {tid: "كلمة " * 300}, None, lib=real)
    except TextDoesNotFit as exc:
        assert exc.details[0]["verdict"] in ("does_not_fit", "missing_glyphs")
    else:  # the template font is not installed here: unmeasured, reported, not refused
        assert rows[0]["verdict"] == "font_missing" and rows[0]["unmeasured"] is True


def test_real_mxlibs_manifests_verify_for_vendoring(monkeypatch):
    if not (REPO / "mxlibs" / "mx_textfit" / "MANIFEST.json").is_file():
        pytest.skip("mxlibs not built yet")
    monkeypatch.delenv("MXS_MXLIBS_DIR", raising=False)
    for lib in ("mx_textfit", "mx_preflight2"):
        try:
            files = lib_files(lib)
        except PatchError as exc:  # a sibling library mid-edit: reported, never vendored
            pytest.skip(f"{lib} manifest not current: {exc}")
        assert "MANIFEST.json" in files and "__init__.py" in files


def test_preflight2_wrapper_shapes():
    calls = []

    def run(path, *, magazine_size="A4", bleed_mm=0):
        calls.append((path, magazine_size, bleed_mm))
        return {"ok": True, "checks": [{"id": "is_pdf", "ok": True}]}

    lib = SimpleNamespace(run=lambda p, profile: {"ok": False, "warnings": ["w"],
                                                   "checks": [{"id": "bleed", "ok": False, "profile": profile}]})
    wrapped = wrap_run(run, lib)
    res = wrapped("x.pdf", magazine_size="A5")
    assert calls == [("x.pdf", "A5", 0)] and wrapped.__mag_wrapped__ is run
    assert res["ok"] is False and res["checks"][-1]["source"] == "mx_preflight2"
    assert res["checks"][-1]["profile"] == {"size": "A5", "bleed_mm": 0.0, "expected_pages": None}
    assert res["mag_preflight2"]["warnings"] == ["w"]
    obj = SimpleNamespace(checks=[], ok=True)
    assert append_checks(obj, {"ok": False, "checks": [{"id": "fonts_embedded", "ok": False}]}).ok is False
    broken = wrap_run(run, SimpleNamespace(run=lambda *a: 1 / 0))("y.pdf")
    assert broken["checks"][-1]["id"] == "mag_preflight2_error" and broken["ok"] is False
    assert profile_from({"max_bytes": 10})["max_bytes"] == 10


def test_safezone_geometry_reports_frames_crossing_the_safe_band(replica):
    g = geometry(_deck(replica()), safe_mm=5)
    assert g["trim_mm"] == [148.0, 210.0]
    assert g["safe_emu"][0] == 5 * 36000
    assert [s["slide"] for s in g["slides"]] == [1, 2, 3]
    tight = geometry(_deck(replica()), safe_mm=40)  # negative control: a huge safe band flags the text boxes
    assert any(f["kind"] == "text" for s in tight["slides"] for f in s["findings"])


def test_single_slide_preview_copy_keeps_one_slide_and_its_slots(replica, tmp_path):
    deck = _deck(replica())
    out = tmp_path / "one.pptx"
    single_slide_copy(deck, 1, out)
    with zipfile.ZipFile(out) as z:
        slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    assert len(slides) == 1 and deck_picture_slots(out) == {1: 1}
    hidden = tmp_path / "hidden.pptx"
    single_slide_copy(deck, 3, hidden)
    assert b'show="0"' not in zipfile.ZipFile(hidden).read("ppt/slides/slide3.xml")
    with pytest.raises(ValueError):
        single_slide_copy(deck, 9, tmp_path / "x.pptx")
    exported = []

    def export_pdf(src, pdf, expect_pages, attempts=2):
        exported.append(expect_pages)
        Path(pdf).write_bytes(b"%PDF-1.4 one page")

    first = cached_preview(deck, 1, tmp_path / "cache", export_pdf)
    second = cached_preview(deck, 1, tmp_path / "cache", export_pdf)
    assert first == second and exported == [1]  # cached by content hash
    assert sorted(p.suffix for p in (tmp_path / "cache").iterdir()) == [".pdf"]


def test_proof_approval_is_bound_to_pdf_bytes(tmp_path):
    pdf = tmp_path / "order_final.pdf"
    pdf.write_bytes(b"%PDF-1.4 v1")
    assert status(pdf)["approved"] is False
    rec = record(pdf, version="7", order_ref="R1")
    assert rec["bound_via"] == "sidecar" and len(rec["sha256"]) == 64
    assert status(pdf)["valid"] is True
    pdf.write_bytes(b"%PDF-1.4 v2 after a customer edit")
    after = status(pdf)
    assert after["approved"] and after["valid"] is False and "changed" in after["reason"]
    bound = []
    service = SimpleNamespace(record_proof_approval=lambda **kw: bound.append(kw))
    assert record(pdf, version="8", order_ref="R1", print_service=service)["bound_via"] == \
        "print_service.record_proof_approval"
    assert bound[0]["order_ref"] == "R1" and bound[0]["actor"] == "customer"


def test_payloads_import_nothing_from_mxstudio():
    base = REPO / "mxstudio" / "kit" / "patches" / "payloads"
    for path in base.rglob("*.py"):
        text = path.read_text(encoding="utf-8")
        assert "from mxstudio" not in text and "import mxstudio" not in text, path
    assert "subprocess" not in "".join(p.read_text(encoding="utf-8") for p in base.rglob("*.py"))
    assert sys.modules.get("mxstudio.kit.patches.payloads.textfit.gate") is not None
