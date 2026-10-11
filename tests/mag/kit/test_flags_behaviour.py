"""Flags OFF ⇒ the hooked functions behave byte-identically; flags ON ⇒ each extension does its job.

Patched replica modules run in a fresh interpreter (the replica's ``app`` package must not leak
between tests); results are compared with an unpatched twin replica.
"""
from __future__ import annotations

import pytest
from kit_helpers import ok_runner, run_child

from mxstudio.kit import apply as kit
from mxstudio.kit.patches import PATCHES

OM_PROBE = r'''
import hashlib, json, os
from pathlib import Path
from pptx import Presentation
from app.editor import deck
from app.printing import preflight
out = {}
order = Path("orders/ORD-REPLICA-1")
src = order / "deck.pptx"
inv = deck.inventory(src)
tid = sorted(k for k, v in inv["texts"].items() if v)[0]
dst = Path(os.environ["MX_OUT"])
try:
    res = deck.apply(src, dst, {tid: os.environ.get("MX_NEW_TEXT", "نص قصير")}, {}, inv)
    prs = Presentation(str(dst))
    out["apply"] = {"ok": True, "result": res, "texts": [p.text for s in prs.slides for sh in s.shapes
                                                         if sh.has_text_frame for p in sh.text_frame.paragraphs]}
except Exception as exc:
    out["apply"] = {"ok": False, "type": type(exc).__name__, "msg": str(exc), "code": getattr(exc, "code", None),
                    "is_value_error": isinstance(exc, ValueError)}
out["preflight"] = preflight.run(order / "order_final.pdf", magazine_size="A5")
out["preflight_code"] = hashlib.sha256(preflight.run.__code__.co_code).hexdigest()
out["preflight_wrapped"] = hasattr(preflight.run, "__mag_wrapped__")
from fastapi.testclient import TestClient
from app.main import app
c = TestClient(app)
paths = ["/api/editor/mag/safezone/ORD-REPLICA-1", "/api/editor/mag/preview/ORD-REPLICA-1/1",
         "/api/editor/mag/proof/ORD-REPLICA-1", "/internal/mag/v1/health"]
out["routes"] = {p: c.get(p).status_code for p in paths}
page = c.get("/editor/ORD-REPLICA-1").text
out["page_scripts"] = [s for s in ("safezone.js", "proof.js") if s in page]
out["route_count"] = len(app.routes)
if os.environ.get("MODULEX_MAG_SAFEZONE") == "1":
    out["safezone"] = c.get(paths[0]).json()
if os.environ.get("MODULEX_MAG_PAGE_PREVIEW") == "1":
    r = c.get(paths[1])
    out["preview"] = {"status": r.status_code, "pdf": r.content[:5].decode("latin-1"),
                      "cached": len(list(Path("mag_cache/preview").glob("*.pdf")))}
if os.environ.get("MODULEX_MAG_PROOF_APPROVAL") == "1":
    out["proof_before"] = c.get(paths[2]).json()
    out["proof_approve"] = c.post(paths[2] + "/approve", json={"version": "v1"}).json()
    out["proof_after"] = c.get(paths[2]).json()
if os.environ.get("MODULEX_MAG_RECEIVERS") == "1":
    from modulex_core import modulex_signing as s
    h = s.sign(os.environ["MODULEX_HMAC_SECRET"], "GET", paths[3], b"")
    out["signed_health"] = c.get(paths[3], headers=h).status_code
print(json.dumps(out, ensure_ascii=False, default=str))
'''

CONSOLE_PROBE = r'''
import json, os
from fastapi.testclient import TestClient
from execution_console.server import app
c = TestClient(app)
out = {"unsigned": c.get("/internal/mag/v1/health").status_code, "routes": len(app.routes)}
if os.environ.get("MODULEX_MAG_RECEIVERS") == "1":
    from execution_console import modulex_signing as s
    h = s.sign(os.environ["MODULEX_HMAC_SECRET"], "GET", "/internal/mag/v1/health", b"")
    out["signed"] = c.get("/internal/mag/v1/health", headers=h).json()
print(json.dumps(out, ensure_ascii=False))
'''

ALL_FLAGS_OFF = {p.flag: "0" for p in PATCHES.values()}


@pytest.fixture()
def twins(replica):
    """(pristine, patched-with-all-six) replicas with identical starting bytes."""
    pristine, patched = replica(), replica()
    root = patched["root"]
    for pid in PATCHES:
        code = next(r for r in kit.plan(root)["patches"] if r["id"] == pid)["confirm_code"]
        assert kit.apply_patch(pid, root=root, confirm=code, test_runner=ok_runner)["result"] == "applied"
    return pristine, patched


def _om(facts, tmp_path, name, **env):
    return run_child(OM_PROBE, pythonpath=facts["om"], env={"MX_OUT": str(tmp_path / f"{name}.pptx"),
                                                            **ALL_FLAGS_OFF, **env})


def test_flags_off_behaviour_is_identical(twins, tmp_path):
    pristine, patched = twins
    a = _om(pristine, tmp_path, "a")
    b = _om(patched, tmp_path, "b")
    assert a == b
    assert a["apply"]["ok"] and a["preflight_wrapped"] is False
    assert set(a["routes"].values()) == {404} and a["page_scripts"] == []
    long_text = "كلمة " * 150
    assert _om(pristine, tmp_path, "c", MX_NEW_TEXT=long_text) == _om(patched, tmp_path, "d", MX_NEW_TEXT=long_text)
    ca = run_child(CONSOLE_PROBE, pythonpath=pristine["console"], env=ALL_FLAGS_OFF)
    cb = run_child(CONSOLE_PROBE, pythonpath=patched["console"], env=ALL_FLAGS_OFF)
    assert ca == cb == {"unsigned": 404, "routes": ca["routes"]}


def test_negative_control_flag_on_changes_behaviour(twins, tmp_path):
    """Proves the comparison above can fail: with the flags on, the patched twin differs."""
    pristine, patched = twins
    on = {"MODULEX_MAG_SAFEZONE": "1", "MODULEX_MAG_PREFLIGHT2": "1"}
    assert _om(pristine, tmp_path, "a", **on) != _om(patched, tmp_path, "b", **on)


def test_textfit_on_refuses_overflow_in_arabic_and_passes_short_text(twins, tmp_path):
    _, patched = twins
    ok = _om(patched, tmp_path, "ok", MODULEX_MAG_TEXTFIT="1")
    assert ok["apply"]["ok"]
    bad = _om(patched, tmp_path, "bad", MODULEX_MAG_TEXTFIT="1", MX_NEW_TEXT="كلمة " * 150)
    assert bad["apply"]["ok"] is False and bad["apply"]["code"] == "MAG_TEXT_OVERFLOW"
    assert bad["apply"]["is_value_error"] and "من فضلك اختصره" in bad["apply"]["msg"]
    assert not (tmp_path / "bad.pptx").exists()  # refused before the deck was written


def test_preflight2_on_appends_named_checks(twins, tmp_path):
    _, patched = twins
    res = _om(patched, tmp_path, "p", MODULEX_MAG_PREFLIGHT2="1")
    assert res["preflight_wrapped"] is True
    ids = [(c["id"], c.get("source")) for c in res["preflight"]["checks"]]
    assert ("bleed", "mx_preflight2") in ids and ("is_pdf", None) in ids
    assert res["preflight"]["ok"] is True  # OM default bleed 0 mm is passed through to the profile


def test_editor_extensions_mount_only_with_their_flags(twins, tmp_path):
    _, patched = twins
    res = _om(patched, tmp_path, "e", MODULEX_MAG_SAFEZONE="1", MODULEX_MAG_PAGE_PREVIEW="1",
              MODULEX_MAG_PROOF_APPROVAL="1")
    assert res["routes"]["/api/editor/mag/safezone/ORD-REPLICA-1"] == 200
    assert res["routes"]["/internal/mag/v1/health"] == 404  # receivers flag still off
    assert sorted(res["page_scripts"]) == ["proof.js", "safezone.js"]
    sz = res["safezone"]
    assert sz["trim_mm"] == [148.0, 210.0] and sz["safe_mm"] == 5.0
    assert res["preview"] == {"status": 200, "pdf": "%PDF-", "cached": 1}
    assert res["proof_before"]["approved"] is False
    assert res["proof_approve"]["approval"]["bound_via"] == "sidecar"
    assert res["proof_after"]["approved"] and res["proof_after"]["valid"]
    only_safezone = _om(patched, tmp_path, "f", MODULEX_MAG_SAFEZONE="1")
    assert only_safezone["routes"]["/api/editor/mag/proof/ORD-REPLICA-1"] == 404
    assert only_safezone["page_scripts"] == ["safezone.js"]


def test_receivers_mount_in_both_systems_with_signed_health(twins, tmp_path):
    _, patched = twins
    env = {"MODULEX_MAG_RECEIVERS": "1", "MODULEX_HMAC_SECRET": "replica-secret-123"}
    om = _om(patched, tmp_path, "r", **env)
    assert om["routes"]["/internal/mag/v1/health"] == 401 and om["signed_health"] == 200
    con = run_child(CONSOLE_PROBE, pythonpath=patched["console"], env={**ALL_FLAGS_OFF, **env})
    assert con["unsigned"] == 401 and con["signed"]["kinds"] == ["m2"] and con["signed"]["system"] == "Console"
