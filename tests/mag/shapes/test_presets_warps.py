"""Gate A: every ECMA-376 preset geometry (187) and text warp (41) is enumerated, with handles."""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from lxml import etree

from mxstudio.shapes import presets, warps

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
LO_ADJ = Path("/usr/lib/libreoffice/share/filter/oox-drawingml-adj-names")
LO_CS = Path("/usr/lib/libreoffice/share/filter/oox-drawingml-cs-presets")

BRIEF_EXAMPLES = ["cloud", "cloudCallout", "wedgeRoundRectCallout", "wedgeEllipseCallout", "irregularSeal1",
                  "irregularSeal2", "star4", "star32", "heart", "sun", "moon", "lightningBolt", "smileyFace", "wave",
                  "doubleWave", "ribbon", "ribbon2", "ellipseRibbon", "horizontalScroll", "verticalScroll", "plaque",
                  "teardrop", "flowChartProcess", "flowChartOffpageConnector"]
BRIEF_WARPS = ["textArchUp", "textWave1", "textInflate", "textCircle"]


def diff_names(names, reference) -> dict:
    names, reference = list(names), set(reference)
    return {"missing": sorted(reference - set(names)), "extra": sorted(set(names) - reference),
            "duplicates": sorted({n for n in names if names.count(n) > 1})}


def test_preset_enumeration_is_the_full_spec_list():
    names = presets.ECMA_PRESET_SHAPES
    assert len(names) == 187 and len(set(names)) == 187
    ref = json.loads((HERE / "fixtures" / "preset_reference.json").read_text(encoding="utf-8"))["presets"]
    assert diff_names(names, ref) == {"missing": [], "extra": [], "duplicates": []}
    for name in BRIEF_EXAMPLES:
        assert name in names
    assert {f"star{n}" for n in (4, 5, 6, 7, 8, 10, 12, 16, 24, 32)} <= set(names)
    assert sum(n.startswith("flowChart") for n in names) == 29


def test_python_pptx_prst_names_are_a_subset():
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

    pptx_names = set()
    for member in MSO_AUTO_SHAPE_TYPE:
        try:
            pptx_names.add(MSO_AUTO_SHAPE_TYPE.to_xml(member))
        except (KeyError, ValueError):
            continue
    assert len(pptx_names) >= 170
    assert pptx_names <= set(presets.ECMA_PRESET_SHAPES)


def test_adjust_handles_match_preset_definitions():
    ref = json.loads((HERE / "fixtures" / "preset_reference.json").read_text(encoding="utf-8"))["presets"]
    mismatches = {p.prst: ([(h.name, h.default) for h in p.handles], ref[p.prst]) for p in presets.all_presets()
                  if [[h.name, h.default] for h in p.handles] != ref[p.prst]}
    assert mismatches == {}
    total = sum(len(p.handles) for p in presets.all_presets())
    assert total == sum(len(v) for v in ref.values()) == 300


@pytest.mark.skipif(not (LO_ADJ.is_file() and LO_CS.is_file()), reason="LibreOffice preset tables not installed")
def test_live_libreoffice_tables_agree():
    lo_names = set(re.findall(r"^/\* ([A-Za-z0-9]+) \*/", LO_CS.read_text(encoding="utf-8"), re.M))
    assert diff_names(presets.ECMA_PRESET_SHAPES, lo_names)["missing"] == []
    lo_adj: dict[str, list[str]] = {}
    for line in LO_ADJ.read_text(encoding="utf-8").splitlines():
        shape, adj = line.split("\t")
        lo_adj.setdefault(shape, []).append(adj)
    for p in presets.all_presets():
        assert [h.name for h in p.handles] == lo_adj.get(p.prst, []), p.prst
    lo_warp = {k: v for k, v in lo_adj.items() if k.startswith("text")}
    for w in warps.all_warps():
        assert [h[0] for h in w.handles] == lo_warp.get(w.prst, []), w.prst


def test_negative_control_enumeration_check_detects_drift():
    corrupted = [n for n in presets.ECMA_PRESET_SHAPES if n != "cloud"] + ["cloudy", "star5"]
    ref = json.loads((HERE / "fixtures" / "preset_reference.json").read_text(encoding="utf-8"))["presets"]
    d = diff_names(corrupted, ref)
    assert d["missing"] == ["cloud"] and d["extra"] == ["cloudy"] and d["duplicates"] == ["star5"]


def test_preset_xml_and_handle_rules():
    xml = presets.prst_geom_xml("wedgeRoundRectCallout", {"adj1": -40000, "adj2": 90000}, with_ns=True)
    el = etree.fromstring(xml)
    gds = {g.get("name"): g.get("fmla") for g in el.iter("{*}gd")}
    assert el.get("prst") == "wedgeRoundRectCallout"
    assert gds == {"adj1": "val -40000", "adj2": "val 90000", "adj3": "val 16667"}
    with pytest.raises(ValueError):
        presets.resolve_adjustments("cloud", {"adj9": 1})
    with pytest.raises(KeyError):
        presets.get_preset("notAShape")
    blockarc = presets.get_preset("blockArc")
    assert [h.kind for h in blockarc.handles] == ["angle", "angle", "ratio"]


def test_warp_enumeration_and_defaults():
    names = warps.ECMA_TEXT_WARPS
    assert len(names) == 41 and len(set(names)) == 41
    for name in BRIEF_WARPS:
        assert name in names
    ref = json.loads((HERE / "fixtures" / "text_warp_reference.json").read_text(encoding="utf-8"))["warps"]
    assert diff_names(names, set(ref) | {"textNoShape"}) == {"missing": [], "extra": [], "duplicates": []}
    for w in warps.all_warps():
        if w.prst == "textNoShape":
            assert w.handles == ()
            continue
        expected = [[n, f.replace("cd2", "10800000")] for n, f in ref[w.prst]]
        assert [[n, f"val {d}"] for n, d, *_ in w.handles] == expected, w.prst


def test_warp_xml_and_negative_control():
    xml = warps.prst_tx_warp_xml("textWave1", {"adj1": 20000}, with_ns=True)
    el = etree.fromstring(xml)
    assert el.get("prst") == "textWave1"
    assert [g.get("fmla") for g in el.iter("{*}gd")] == ["val 20000", "val 0"]
    with pytest.raises(ValueError):
        warps.resolve("textWave1", {"adj": 1})
    with pytest.raises(KeyError):
        warps.get_warp("textSpiral")


def test_catalog_snapshots_match_code():
    from mxstudio.shapes.generators import list_generators

    cat = ROOT / "catalog" / "shapes"
    assert json.loads((cat / "presets.json").read_text(encoding="utf-8")) == json.loads(json.dumps(presets.catalog()))
    assert json.loads((cat / "warps.json").read_text(encoding="utf-8")) == json.loads(json.dumps(warps.catalog()))
    assert json.loads((cat / "generators.json").read_text(encoding="utf-8")) == json.loads(
        json.dumps({"generators": list_generators()}))
