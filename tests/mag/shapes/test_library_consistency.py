"""Gate F: library (tags, favourites, usage, text + sketch search) and book style consistency."""
from __future__ import annotations

import cv2
import numpy as np
import pytest

from mxstudio.shapes import consistency
from mxstudio.shapes.custgeom import CustomGeometry, GeomPart
from mxstudio.shapes.discovery.base import Provenance
from mxstudio.shapes.generators import generate
from mxstudio.shapes.generators._util import circle, polygon, star_points
from mxstudio.shapes.library import LibraryError, fold_text


def geom(path, w=1000, h=1000):
    return CustomGeometry(w, h, [GeomPart(path)])


def seed(library):
    ids = {}
    ids["circle"] = library.add_original(geom(circle(500, 500, 400)), "generator", author="Studio",
                                         source_url="mxstudio://generator/circle", name_ar="دائرة", name_en="circle",
                                         tags=["شمس", "basic"], theme="space", style="flat")
    ids["star"] = library.add_original(geom(polygon(star_points(5, 500, 520, 450, 180))), "generator", author="Studio",
                                       source_url="mxstudio://generator/star", name_ar="نجمة", name_en="star",
                                       tags=["night"], theme="space", style="outline")
    ids["rabbit"] = library.add_original(generate("blob", 3).layers[0].geometry, "canvas", author="Mustafa",
                                         source_url="mxstudio://canvas", name_ar="أرنب", name_en="rabbit blob",
                                         tags=["حيوانات"], theme="farm", age="3-5", style="hand-drawn")
    ids["square"] = library.add_original(geom(polygon([(100, 100), (900, 100), (900, 900), (100, 900)])), "preset",
                                         author="Studio", source_url="mxstudio://preset/rect", name_en="square",
                                         style="flat")
    return ids


def test_text_search_with_arabic_folding_and_filters(library):
    ids = seed(library)
    assert [a.id for a in library.search_text("ارنب")] == [ids["rabbit"]]  # hamza-less spelling still matches
    assert [a.id for a in library.search_text("أَرْنَب")] == [ids["rabbit"]]  # tashkeel ignored
    assert {a.id for a in library.search_text("", theme="space")} == {ids["circle"], ids["star"]}
    assert [a.id for a in library.search_text("st", style="outline")] == [ids["star"]]
    assert library.search_text("zebra") == []
    assert fold_text("مكتبة الأشكالِ") == "مكتبه الاشكال"


def test_favourites_and_usage(library):
    ids = seed(library)
    library.set_favourite(ids["star"])
    assert [a.id for a in library.favourites()] == [ids["star"]]
    assert library.record_use(ids["circle"], "book-7", 3) == 1
    assert library.record_use(ids["circle"], "book-7", 4) == 2
    library.record_use(ids["star"], "book-7", 4)
    assert [a.id for a in library.most_used()] == [ids["circle"], ids["star"]]
    assert len(library.uses_in_book("book-7")) == 3
    library.set_tags(ids["square"], ["box", "مربع"])
    assert [a.id for a in library.search_text("مربع")] == [ids["square"]]
    with pytest.raises(KeyError):
        library.record_use("missing", "book-7")


def sketch(kind: str) -> np.ndarray:
    img = np.zeros((300, 300), np.uint8)
    if kind == "circle":
        cv2.ellipse(img, (150, 150), (110, 104), 10, 0, 360, 1, -1)
    elif kind == "star":
        pts = np.array(star_points(5, 150, 160, 130, 55, -80), np.int32)
        cv2.fillPoly(img, [pts], 1)
    return img


def test_sketch_similarity_ranks_the_right_shape_first(library):
    ids = seed(library)
    assert library.search_similar(mask=sketch("circle"))[0][1].id == ids["circle"]
    hits = library.search_similar(mask=sketch("star"))
    assert hits[0][1].id == ids["star"]
    # negative control: the circle sketch is measurably farther from the star than from the circle
    circle_hits = {a.id: d for d, a in library.search_similar(mask=sketch("circle"))}
    assert circle_hits[ids["circle"]] < circle_hits[ids["star"]]
    rotated = library.search_similar(geometry=geom(polygon(star_points(5, 500, 500, 450, 180, -50))))
    assert rotated[0][1].id == ids["star"]  # rotation-invariant


def test_admission_rules(library):
    g = geom(circle(500, 500, 300))
    good = Provenance(source="openverse", source_url="https://x.example/1", author="Ann", licence="CC-BY-4.0",
                      licence_url="https://creativecommons.org/licenses/by/4.0/", retrieved_at="2026-10-11",
                      sha256="e" * 64, credits_required=True)
    with pytest.raises(LibraryError, match="approval"):
        library.add(geometry=g, origin="discovered", provenance=good)
    asset_id = library.add(geometry=g, origin="discovered", provenance=good, approved_by="Mustafa")
    assert library.get(asset_id).credits_required
    bad = Provenance(**{**good.to_json(), "licence": "CC-BY-SA-4.0"})
    with pytest.raises(LibraryError, match="licence refused"):
        library.add(geometry=g, origin="discovered", provenance=bad, approved_by="Mustafa")
    incomplete = Provenance(**{**good.to_json(), "sha256": ""})
    with pytest.raises(LibraryError, match="provenance incomplete"):
        library.add(geometry=g, origin="discovered", provenance=incomplete, approved_by="Mustafa")
    with pytest.raises(LibraryError, match="style"):
        library.add_original(g, "canvas", author="M", source_url="mxstudio://canvas", style="watercolour")


def test_consistency_findings_and_clean_book():
    from pptx import Presentation  # noqa: F401  (python-pptx is the deck reader's dependency)

    uses = [consistency.ShapeUse("1", "a", 2.0, 16667, ["#E63946"], "flat"),
            consistency.ShapeUse("2", "b", 2.1, 16000, ["#E63946", "#FFFFFF"], "flat"),
            consistency.ShapeUse("3", "c", 6.0, 45000, ["#00FF66"], "hand-drawn")]
    out = consistency.check(uses, palette=["#E63946", "#457B9D", "#F4A261"])
    codes = sorted({f["code"] for f in out["findings"]})
    assert codes == ["corner_radius_outlier", "mixed_styles", "off_palette_colour", "stroke_width_outlier"]
    assert [f["page"] for f in out["findings"] if f["code"] == "stroke_width_outlier"] == ["3"]
    clean = consistency.check(uses[:2], palette=["#E63946"])
    assert clean["findings"] == []  # negative control: a consistent book raises nothing


def test_consistency_reads_a_deck(tmp_path):
    from mxstudio.shapes import pptx_out

    prs = pptx_out.new_presentation()
    s1 = pptx_out.blank_slide(prs)
    pptx_out.add_preset(s1.shapes, "roundRect", 100000, 100000, 2000000, 1000000, fill="#E63946", line="#222222",
                        line_w_pt=2)
    s2 = pptx_out.blank_slide(prs)
    pptx_out.add_preset(s2.shapes, "roundRect", 100000, 100000, 2000000, 1000000, adj={"adj": 45000}, fill="#E63946",
                        line="#222222", line_w_pt=2)
    pptx_out.add_custgeom(s2.shapes, geom(circle(500, 500, 400)), 100000, 1300000, 1000000, 1000000, fill="#13FF13",
                          line="#222222", line_w_pt=7)
    path = tmp_path / "book.pptx"
    prs.save(path)
    uses = consistency.uses_from_pptx(path)
    assert len(uses) == 3 and uses[0].corner_radius == 16667.0 and uses[1].corner_radius == 45000.0
    out = consistency.check(uses, palette=["#E63946"])
    codes = {f["code"] for f in out["findings"]}
    assert {"stroke_width_outlier", "off_palette_colour", "corner_radius_outlier"} <= codes
