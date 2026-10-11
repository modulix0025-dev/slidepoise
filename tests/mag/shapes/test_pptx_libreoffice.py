"""Gate: a PPTX with every preset, every warp, every generator and converted SVGs is valid XML,
stays native (sp + custGeom/prstGeom, no pictures), and LibreOffice renders it to PDF.

The arcTo check also proves LibreOffice reads ``arcTo`` angles the way this package writes them
(visual angles): its PDF render of arc-heavy corpus shapes overlaps resvg's render of the SVG.
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import numpy as np
import pytest
from lxml import etree

ROOT = Path(__file__).resolve().parents[3]
CORPUS = Path(__file__).resolve().parent / "corpus"
sys.path.insert(0, str(ROOT / "vendor" / "modulex"))
import pptx_to_pdf  # noqa: E402  (vendored, byte-identical)

from mxstudio.shapes import pptx_out  # noqa: E402
from mxstudio.shapes.raster import iou, render_svg_mask  # noqa: E402
from mxstudio.shapes.svg_convert import svg_to_shapes  # noqa: E402

ARC_SHAPES = ["18_arc_elliptic", "19_arc_elliptic_offaxis", "21_arc_rotated_90", "34_nonzero_donut_opposite",
              "38_heart", "40_moon"]


@pytest.fixture(scope="module")
def proof_deck(tmp_path_factory):
    out = tmp_path_factory.mktemp("proof") / "shapes_proof.pptx"
    svgs = [CORPUS / f"{n}.svg" for n in ("35_multi_element", "46_letter_A_like", "50_mixed_everything")]
    info = pptx_out.build_proof_deck(out, svg_files=svgs)
    return out, info


def test_deck_xml_is_valid_and_native(proof_deck):
    path, info = proof_deck
    assert info["findings"] == []
    counts = info["counts"]
    assert counts["pic"] == 0 and counts["prstTxWarp"] == 41 and counts["custGeom"] >= 25
    with zipfile.ZipFile(path) as z:
        prst = set()
        for name in z.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                root = etree.fromstring(z.read(name))
                prst |= {e.get("prst") for e in root.iter("{*}prstGeom")}
                for cg in root.iter("{*}custGeom"):
                    assert etree.QName(cg.getparent().getparent()).localname == "sp"
    from mxstudio.shapes.presets import ECMA_PRESET_SHAPES

    assert set(ECMA_PRESET_SHAPES) <= prst


def test_deck_reopens_and_shapes_are_editable(proof_deck):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    path, info = proof_deck
    prs = Presentation(str(path))
    assert len(prs.slides) == info["slides"]
    freeforms = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE
            items = shape.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP else [shape]
            for sh in items:
                if sh._element.find(".//{*}custGeom") is not None:
                    freeforms += 1
                    sh.fill.solid()  # editable: recolouring works through the normal API
    assert freeforms == info["counts"]["custGeom"]


def test_libreoffice_renders_the_deck(proof_deck, soffice, tmp_path):
    path, info = proof_deck
    rec = pptx_to_pdf.convert(path, tmp_path / "shapes_proof.pdf", engines=["soffice"], expect_pages=info["slides"])
    assert rec["ok"] and rec["pages"] == info["slides"] and "page_mismatch" not in rec


def test_libreoffice_reads_arcto_like_we_write_it(soffice, tmp_path):
    import pypdfium2 as pdfium
    from pptx.util import Pt

    prs = pptx_out.new_presentation(100, 100)
    for name in ARC_SHAPES:
        res = svg_to_shapes((CORPUS / f"{name}.svg").read_text(encoding="utf-8"))
        slide = pptx_out.blank_slide(prs)
        tb = slide.shapes.add_textbox(0, 0, pptx_out.mm(8), pptx_out.mm(4))
        tb.text_frame.text = "."
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(4)
        k = pptx_out.mm(100) / res.width
        for shape in res.shapes:
            local, (x, y, w, h) = shape.geometry.cropped()
            pptx_out.add_custgeom(slide.shapes, local, int(x * k), int(y * k), int(w * k), int(h * k), fill="#000000")
    deck = tmp_path / "arcs.pptx"
    prs.save(deck)
    pptx_to_pdf.convert(deck, tmp_path / "arcs.pdf", engines=["soffice"])
    pdf = pdfium.PdfDocument(str(tmp_path / "arcs.pdf"))
    import cv2

    scores = {}
    for i, name in enumerate(ARC_SHAPES):
        page = pdf[i]
        img = np.array(page.render(scale=600 / page.get_width()).to_pil().convert("L"))
        img = cv2.resize(img, (600, 600), interpolation=cv2.INTER_AREA)
        got = (img < 128).astype(np.uint8)
        got[:40, :100] = 0  # the tiny text marker
        ref = render_svg_mask((CORPUS / f"{name}.svg").read_text(encoding="utf-8"), 600, 600)
        scores[name] = iou(got, ref)
    assert min(scores.values()) >= 0.98, scores
