"""Phase 9 gate: every blocker holds the order with its own name; advisory checks never change state."""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from mx_qagate import advisory_review, qa_sheet, run_gate
from mxstudio.factory.demo_story import moon_star_story
from mxstudio.factory.forge import forge
from mxstudio.validators.operator_ports import replay_generate_story

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "vendor" / "modulex"))
import pptx_to_pdf  # noqa: E402


@pytest.fixture(scope="module")
def order(tmp_path_factory, fixture_registry, fixture_fontconfig, soffice, render_deck):
    """A real personalised + injected order: template → OM replacement → scene injection → export."""
    os.environ["FONTCONFIG_FILE"] = str(fixture_fontconfig)
    out = tmp_path_factory.mktemp("order")
    built = forge(moon_star_story(), out, sizes=("A5",), registry=fixture_registry)
    template = Path(built["variants"]["ar_boy_A5"]["path"])
    personalised = out / "محمد - moon_star (AR) [A5].pptx"
    replay_generate_story(template, "إياد", "محمد", out_path=personalised)
    scene = out / "scene.png"
    import numpy as np
    gradient = np.linspace(40, 220, 900, dtype=np.uint8)
    Image.fromarray(np.dstack([np.tile(gradient, (700, 1))] * 3)).save(scene)
    prs = Presentation(personalised)
    for slide in prs.slides:
        for placeholder in slide.placeholders:
            placeholder.insert_picture(str(scene))
    injected = out / "محمد - moon_star (AR) [A5] — ModuleX.pptx"
    prs.save(injected)
    pdf = out / "order_test_final.pdf"
    pptx_to_pdf.convert(injected, pdf, engines=["soffice"])
    return {"dir": out, "personalised": personalised, "injected": injected, "pdf": pdf, "registry": fixture_registry}


def _codes(result):
    return [f["code"] for f in result["findings"]]


def test_a_good_order_passes(order):
    result = run_gate(order["injected"], order["pdf"], approved_pages=6, child_name="محمد",
                      certified_fonts={"Tajawal"}, installed_fonts=set(order["registry"].families()))
    assert result["hold"] is False, result["findings"]
    assert result["facts"]["pages"] == 6 and result["facts"]["blank_check"] == "rendered"


def test_empty_placeholder_holds(order, tmp_path):
    pdf = tmp_path / "uninjected.pdf"
    pptx_to_pdf.convert(order["personalised"], pdf, engines=["soffice"])
    result = run_gate(order["personalised"], pdf, approved_pages=6, child_name="محمد")
    assert "empty_placeholder" in _codes(result) and result["hold"]
    assert "blank_page" in _codes(result)  # the back page has only (empty) placeholders


def test_page_count_and_wrong_child_name_hold(order):
    result = run_gate(order["injected"], order["pdf"], approved_pages=16, child_name="عمر")
    assert {"page_count", "child_name_missing"} <= set(_codes(result))
    for finding in result["findings"]:
        assert finding["ar"] and finding["en"] and finding["blocking"] is True


def test_image_only_pdf_holds(order, tmp_path):
    pages = [Image.new("RGB", (437, 612), (200, 200, 255)) for _ in range(6)]
    for page in pages:
        page.putpixel((10, 10), (0, 0, 0))
    pdf = tmp_path / "pictures.pdf"
    pages[0].save(pdf, save_all=True, append_images=pages[1:])
    result = run_gate(order["injected"], pdf, approved_pages=6, child_name="محمد", check_overflow=False)
    assert "image_only_pdf" in _codes(result)


def test_uncertified_font_holds(order):
    result = run_gate(order["injected"], order["pdf"], approved_pages=6, child_name="محمد",
                      certified_fonts={"Amiri"}, check_overflow=False)
    assert _codes(result) == ["font_not_certified"]
    assert result["findings"][0]["detail"]["fonts"] == ["Tajawal"]


def test_text_overflow_holds(order, tmp_path):
    prs = Presentation(order["injected"])
    body = next(s for s in prs.slides[1].shapes if s.has_text_frame and s.text_frame.paragraphs[0].runs)
    body.text_frame.paragraphs[0].runs[0].text = "كان يا ما كان " * 80
    deck = tmp_path / "long.pptx"
    prs.save(deck)
    from slidepoise.arabic.fit import fit_text

    def fit(text, w_pt, h_pt, family, size_pt):
        return fit_text(text, w_pt, h_pt, family, registry=order["registry"], max_size_pt=size_pt,
                        min_size_pt=size_pt, inset_pt=0.0).fits

    assert run_gate(order["injected"], order["pdf"], approved_pages=6, child_name="محمد", fit=fit)["hold"] is False
    result = run_gate(deck, order["pdf"], approved_pages=6, child_name="محمد", fit=fit)
    assert "text_overflow" in _codes(result)


def test_advisory_never_changes_state(order, tmp_path):
    gate = run_gate(order["injected"], order["pdf"], approved_pages=6, child_name="محمد")
    worried = advisory_review(lambda check, page: {"concern": True, "reason": "looks off"}, [{"page": 1}, {"page": 2}])
    assert all(item["advisory"] and item["concern"] for item in worried)
    assert gate["hold"] is False
    record = qa_sheet(gate, worried, tmp_path / "order_test_final.pdf.qa.json")
    import json
    data = json.loads(record.read_text(encoding="utf-8"))
    assert data["hold"] is False and data["advisory_changes_state"] is False and len(data["advisory"]) == 8
    assert advisory_review(None, [{"page": 1}])[0]["status"] == "not_configured"
