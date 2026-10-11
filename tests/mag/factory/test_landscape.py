"""16:9 landscape pages with a full-page scene (phase 16, item 11).

Acceptance: factory build with page=landscape_169 and layout=storybook gives 8 green variants; the empty template's
scene slots are real picture placeholders (idx 100, «Picture Placeholder 100», descr «scene»), empty; python-pptx
insert_picture fills every one. Plus: effects in OOXML, safe zone with rotation, contrast on the real art.
"""
from __future__ import annotations

import io
import os
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from PIL import Image

from mxstudio.factory.demo_story import moon_star_story
from mxstudio.factory.story import Story

ROOT = Path(__file__).resolve().parents[3]
FONTS = ROOT / "tests" / "mag" / "fixtures" / "fonts"
ready = pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("node")
                                and (ROOT / "node_modules" / "pptxgenjs").is_dir()),
                           reason="LibreOffice + Node + pptxgenjs are needed for a real build")


def _story(layout: str, **page_over) -> Story:
    story = moon_star_story()
    story.page, story.layout = "landscape_169", layout
    for page in story.pages:
        for key, value in page_over.items():
            setattr(page, key, value)
    return story


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    base = tmp_path_factory.mktemp("land")
    conf = base / "fonts.conf"
    conf.write_text('<?xml version="1.0"?><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                    f"<dir>{FONTS}</dir><cachedir>{base / 'fc'}</cachedir></fontconfig>", encoding="utf-8")
    os.environ["FONTCONFIG_FILE"] = str(conf)
    os.environ["MODULEX_SOFFICE"] = shutil.which("soffice") or ""
    from slidepoise.arabic.fonts import FontRegistry
    return FontRegistry(directories=[FONTS])


def _build(story, out, registry):
    from mxstudio.factory.service import build_and_validate
    return build_and_validate(story, out, registry=registry, engines=["soffice"])


def test_landscape_page_is_16_9_and_scales_exactly_to_a5():
    from slidepoise.print_geometry import PageSpec
    a4 = PageSpec.preset("A4", page_format="landscape_169")
    a5 = PageSpec.preset("A5", page_format="landscape_169")
    assert a4.trim_mm == (297.0, 167.06) and abs(a4.trim_mm[0] / a4.trim_mm[1] - 16 / 9) < 1e-3
    assert abs(a5.trim_mm[0] - 297 / 2 ** 0.5) < 0.01 and a4.canvas_mm == (303.0, 173.06)


def test_story_fields_are_validated():
    story = _story("storybook")
    assert story.validate() == []
    story.pages[1].text_box = [0.8, 0.1, 0.4, 0.2]  # leaves the page
    story.pages[1].name_style = "silver"
    problems = story.validate()
    assert any("text_box" in p for p in problems) and any("name_style" in p for p in problems)
    assert any("layout" in p for p in Story.from_dict({**_story("x").as_dict(), "layout": "x"}).validate())


@ready
@pytest.mark.parametrize("layout", ["storybook", "overlay"])
def test_acceptance_landscape_builds_eight_green_variants_with_real_empty_placeholders(layout, env, tmp_path):
    result = _build(_story(layout), tmp_path, env)
    variants = result["validation"]["variants"]
    assert len(variants) == 8 and result["validation"]["ok"], {k: [f["code"] for f in v["findings"]]
                                                               for k, v in variants.items() if not v["ok"]}
    from pptx import Presentation
    for key, entry in result["built"]["variants"].items():
        deck = Path(entry["path"])
        prs = Presentation(str(deck))
        cw, ch = entry["spec"]["page"]["canvas_mm"]
        assert abs(prs.slide_width / 36000 - cw) < 0.05 and abs(prs.slide_height / 36000 - ch) < 0.05
        with zipfile.ZipFile(deck) as z:
            for number in range(1, len(prs.slides) + 1):
                xml = z.read(f"ppt/slides/slide{number}.xml").decode("utf-8")
                assert "<p:pic>" not in xml  # never a baked-in picture
                assert re.search(r'name="Picture Placeholder 100" descr="scene"', xml), (key, number)
                assert re.search(r'<p:ph type="pic"[^>]*idx="100"', xml)
            layouts = [z.read(n).decode() for n in z.namelist() if re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n)]
            assert any(re.search(r'<p:ph type="pic"[^>]*idx="100"', x) for x in layouts)
        # python-pptx fills every scene placeholder by its index
        art = io.BytesIO()
        Image.new("RGB", (1920, 1080), "#2E86AB").save(art, "PNG")
        filled = 0
        for slide in prs.slides:
            for ph in slide.placeholders:
                if ph.placeholder_format.idx == 100:
                    art.seek(0)
                    ph.insert_picture(art)
                    filled += 1
        assert filled == len(prs.slides) == len(entry["spec"]["slides"])
        out = tmp_path / f"filled_{key}.pptx"
        prs.save(str(out))
        again = Presentation(str(out))
        assert sum(1 for s in again.slides for sh in s.shapes
                   if sh.is_placeholder and sh.placeholder_format.idx == 100 and sh.shape_type == 14 and
                   hasattr(sh, "image")) == filled


@ready
def test_effects_are_ooxml_powerpoint_renders(env, tmp_path):
    result = _build(_story("storybook", name_style="pink"), tmp_path / "s", env)
    deck = Path(result["built"]["variants"]["ar_girl_A4"]["path"])
    with zipfile.ZipFile(deck) as z:
        cover = z.read("ppt/slides/slide1.xml").decode("utf-8")
        scene = z.read("ppt/slides/slide2.xml").decode("utf-8")
    name_art = cover[cover.index('name="p1_name_art"'):]
    name_art = name_art[:name_art.index("</p:sp>")]
    assert '<a:ln w="38100">' in name_art and "<a:gradFill" in name_art and "<a:glow" in name_art
    assert 'rtl="1"' in name_art
    body = scene[scene.index('name="p2_body"'):]
    assert re.search(r'<a:xfrm rot="21480000"', body[:body.index("</p:sp>")])  # -2 degrees
    overlay = _build(_story("overlay"), tmp_path / "o", env)
    with zipfile.ZipFile(Path(overlay["built"]["variants"]["ar_boy_A4"]["path"])) as z:
        xml = z.read("ppt/slides/slide2.xml").decode("utf-8")
    shade = xml[xml.index('name="p2_shade"'):]
    shade = shade[:shade.index("</p:sp>")]
    assert "<a:gradFill" in shade and '<a:alpha val="78000"/>' in shade and '<a:alpha val="0"/>' in shade
    assert '<a:gs pos="40000">' in shade


@ready
def test_text_outside_the_safe_zone_is_named_even_when_rotated(env, tmp_path):
    story = _story("storybook")
    story.pages[1].text_box = [0.0, 0.0, 0.45, 0.30]  # the operator's own box at the very corner
    story.pages[1].rotate = -8
    result = _build(story, tmp_path, env)
    codes = {f["code"] for v in result["validation"]["variants"].values() for f in v["findings"]}
    assert "text_outside_safe" in codes


def test_contrast_is_measured_on_the_art_under_the_words(tmp_path):
    from mxstudio.factory.forge import build_master
    from mxstudio.validators.landscape import contrast_findings
    bright, dark = tmp_path / "bright.png", tmp_path / "dark.png"
    Image.new("RGB", (1600, 900), "#F4F1E8").save(bright)
    Image.new("RGB", (1600, 900), "#10162B").save(dark)
    story = _story("storybook", art=str(bright))
    low, facts = contrast_findings(build_master(story, "A4"))
    assert low and all(f["suggest"] for f in low) and facts["measured"]
    assert "ink_or_overlay" in low[0]["suggest"] or "move_text_box_to" in low[0]["suggest"]
    assert not contrast_findings(build_master(_story("storybook", art=str(dark)), "A4"))[0]
    # overlay: the dark gradient under the words makes the same bright art readable
    over, _ = contrast_findings(build_master(_story("overlay", art=str(bright)), "A4"))
    assert not over or all("darken_overlay_to_alpha" in f["suggest"] for f in over)
    # no art: nothing is guessed
    none, facts = contrast_findings(build_master(_story("storybook"), "A4"))
    assert not none and facts["unmeasured"]
    # half bright / half dark art: the finding suggests where the box reads
    split = tmp_path / "split.png"
    image = Image.new("RGB", (1600, 900), "#F4F1E8")
    image.paste(Image.new("RGB", (800, 900), "#10162B"), (0, 0))
    image.save(split)
    story = _story("storybook", art=str(split))
    story.pages[1].text_box = [0.55, 0.06, 0.40, 0.30]
    low, _ = contrast_findings(build_master(story, "A4"))
    moved = [f for f in low if f["object"] == "p2_body"]
    assert moved and moved[0]["suggest"]["move_text_box_to"]["x"] < 0.5


@ready
def test_baked_picture_or_missing_placeholder_is_a_finding(env, tmp_path):
    from mxstudio.validators.landscape import scene_placeholder_findings
    result = _build(_story("storybook"), tmp_path, env)
    entry = result["built"]["variants"]["ar_boy_A4"]
    deck = Path(entry["path"])
    assert scene_placeholder_findings(deck, entry["spec"]) == []
    broken = tmp_path / "broken.pptx"
    with zipfile.ZipFile(deck) as src, zipfile.ZipFile(broken, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "ppt/slides/slide2.xml":
                data = data.replace(b'idx="100"', b'idx="7"')
            dst.writestr(info, data)
    assert scene_placeholder_findings(broken, entry["spec"])[0]["slide"] == 2
