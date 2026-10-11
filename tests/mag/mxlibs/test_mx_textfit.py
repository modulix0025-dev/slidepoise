"""mx_textfit: known answers, whitespace-only breaks, exact family names, backend fallback, PPTX boxes."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

from mxlibs.mx_textfit import (
    FontResolver,
    ParagraphNotFound,
    char_limit,
    fit,
    fit_paragraph_edit,
    measurer_for,
    normalize_family,
    paragraph_box,
)
from mxlibs.mx_textfit import measure as tf_measure

FONTS = Path(__file__).resolve().parents[1] / "fixtures" / "fonts"
BODY = "كان يا ما كان، ولد اسمه إياد يحب البحر كثيرًا، ويحلم كل ليلة بسمكة ذهبية تكلّمه."
DEDICATION = "بكل الحب تهدى لـعبد الرحمن"
LONG_DEDICATION = (
    DEDICATION
    + " "
    + ("ونتمنى له حياة سعيدة مليانة فرح وحب ونجاح، وكل سنة وهو طيب وبخير مع بابا وماما وكل العيلة الحلوة دي ") * 4
)
TOO_LONG_AR = "النص ده أطول من المكان المتاح — قصّره شوية"
CM = 360000
A5_DEDICATION_BOX = ((138 * 36000), int(43.2 * 36000))  # safe width x 20 % of the A5 canvas


@pytest.fixture(scope="module")
def resolver():
    return FontResolver(directories=[FONTS])


def tajawal(size=16.0, **extra):
    return {"family": "Tajawal", "size_pt": size, **extra}


# --------------------------------------------------------------------------- known answers
@pytest.mark.parametrize(
    "box_cm,size,lines", [((4, 2), 9.5, 3), ((6, 3), 14.5, 3), ((10, 2.5), 16.5, 2), ((5, 5), 19.0, 5)]
)
def test_known_answer_arabic_sentence_shrinks_to_size(resolver, box_cm, size, lines):
    box = (int(box_cm[0] * CM), int(box_cm[1] * CM))
    result = fit(
        BODY, box, tajawal(24), {"allow_shrink": True, "min_size_pt": 8, "inset_emu": 45720}, resolver=resolver
    )
    assert result["verdict"] == "will_shrink" and result["fits"]
    assert result["size_pt"] == size and result["line_count"] == lines
    assert result["method"] == "harfbuzz" and not result["approximate"]
    assert result["ar"] == f"النص هيصغر لـ {size:g} نقطة عشان يدخل في المكان"


def test_matches_phase1_runtime_fitter_exactly(resolver):
    """Same HarfBuzz rules as slidepoise.arabic.fit_text (the Phase-1 engine), without depending on it."""
    from slidepoise.arabic import FontRegistry, fit_text

    registry = FontRegistry(directories=[FONTS])
    for w_cm, h_cm in ((4, 2), (6, 3), (7.5, 1.6)):
        ours = fit(
            BODY,
            (int(w_cm * CM), int(h_cm * CM)),
            tajawal(24),
            {"allow_shrink": True, "min_size_pt": 8, "inset_emu": 45720},
            resolver=resolver,
        )
        theirs = fit_text(
            BODY,
            w_cm * 72 / 2.54,
            h_cm * 72 / 2.54,
            "Tajawal",
            max_size_pt=24,
            min_size_pt=8,
            inset_pt=3.6,
            registry=registry,
        )
        assert (ours["size_pt"], ours["lines"]) == (theirs.size_pt, theirs.lines)


def test_dedication_fits_at_its_size(resolver):
    result = fit(DEDICATION, A5_DEDICATION_BOX, tajawal(24), resolver=resolver)
    assert result["verdict"] == "fits" and result["size_pt"] == 24 and result["line_count"] == 1
    assert result["ar"] == "النص داخل في المكان تمام"


def test_too_long_dedication_is_refused_in_arabic(resolver):
    result = fit(LONG_DEDICATION, A5_DEDICATION_BOX, tajawal(24), resolver=resolver)
    assert result["verdict"] == "does_not_fit" and not result["fits"]
    assert result["reason"] == "height" and result["overflow_pt"] > 100
    assert result["ar"] == TOO_LONG_AR
    assert result["en"].startswith("This text is longer than the space available")


def test_shrink_floor_still_refuses(resolver):
    result = fit(
        LONG_DEDICATION, A5_DEDICATION_BOX, tajawal(24), {"allow_shrink": True, "min_size_pt": 14}, resolver=resolver
    )
    assert result["verdict"] == "does_not_fit" and result["size_pt"] == 14.0 and result["overflow_pt"] > 0


# --------------------------------------------------------------------------- breaking rules
def test_breaks_only_at_whitespace_and_never_splits_arabic_words(resolver):
    result = fit(BODY, (int(3 * CM), int(10 * CM)), tajawal(16), resolver=resolver)
    assert result["line_count"] > 3
    words = BODY.split()
    assert [w for line in result["lines"] for w in line.split()] == words
    assert all(line == line.strip() for line in result["lines"])


def test_nbsp_does_not_break(resolver):
    glued = "إياد الصغير " * 6
    result = fit(glued, (int(4 * CM), int(10 * CM)), tajawal(14), resolver=resolver)
    assert all("إياد الصغير" in line or not line for line in result["lines"])


def test_word_wider_than_box_is_a_named_overflow(resolver):
    result = fit("عبدالرحمنعبدالرحمنعبدالرحمن", (int(2 * CM), int(3 * CM)), tajawal(18), resolver=resolver)
    assert result["verdict"] == "does_not_fit" and result["reason"] == "word_too_wide"
    assert result["unbreakable_words"] and result["overflow_width_pt"] > 0
    assert "أعرض من المكان" in result["ar"]


def test_explicit_newlines_and_max_lines(resolver):
    text = "\n".join(["سطر"] * 31)
    result = fit(text, (int(10 * CM), int(100 * CM)), tajawal(10), resolver=resolver)
    assert result["line_count"] == 31 and result["reason"] == "too_many_lines"
    assert result["ar"] == "النص ده 31 سطر والمسموح 30 بس — قلّل السطور"
    assert fit("\n".join(["سطر"] * 30), (int(10 * CM), int(100 * CM)), tajawal(10), resolver=resolver)["fits"]


def test_harakat_raise_the_line_height(resolver):
    plain = fit("محمد", (int(10 * CM), int(5 * CM)), {"family": "Amiri", "size_pt": 30}, resolver=resolver)
    marked = fit("مُحَمَّدٌ", (int(10 * CM), int(5 * CM)), {"family": "Amiri", "size_pt": 30}, resolver=resolver)
    assert marked["harakat"] and marked["line_height_pt"] >= plain["line_height_pt"]


def test_editor_char_limit_formula():
    assert char_limit("x" * 10) == 80
    assert char_limit("x" * 500) == 1040
    assert char_limit("x" * 1500) == 2000


def test_max_chars_rule(resolver):
    result = fit("ب" * 90, (int(20 * CM), int(20 * CM)), tajawal(10), {"max_chars": 80}, resolver=resolver)
    assert result["reason"] == "too_many_chars" and not result["fits"]


# --------------------------------------------------------------------------- fonts
def _renamed_font(tmp_path: Path, family: str) -> Path:
    from fontTools.ttLib import TTFont

    font = TTFont(str(FONTS / "tajawal" / "Tajawal-Regular.ttf"))
    for record in font["name"].names:
        if record.nameID in (1, 4, 16):
            record.string = family
    out = tmp_path / "renamed.ttf"
    font.save(str(out))
    return out


def test_exact_family_name_keeps_internal_double_space(tmp_path):
    path = _renamed_font(tmp_path, "Blabeloo  MagdySoliman")
    resolver = FontResolver(directories=[], files=[path])
    assert normalize_family(" Blabeloo  MagdySoliman ") == "blabeloo  magdysoliman"
    assert resolver.resolve("Blabeloo  MagdySoliman") is not None
    assert resolver.resolve("Blabeloo MagdySoliman") is None  # negative control: single space is another name
    result = fit(DEDICATION, A5_DEDICATION_BOX, {"family": "Blabeloo  MagdySoliman", "size_pt": 24}, resolver=resolver)
    assert result["font_used"]["family"] == "Blabeloo  MagdySoliman" and result["substitute"] is None


def test_missing_family_is_reported_not_substituted(resolver):
    result = fit(
        BODY, (int(6 * CM), int(3 * CM)), {"family": "Blabeloo  MagdySoliman", "size_pt": 16}, resolver=resolver
    )
    assert result["verdict"] == "font_missing" and result["font_used"] is None and not result["fits"]
    assert "مش متسطّب" in result["ar"]
    with_fallback = fit(
        BODY,
        (int(6 * CM), int(3 * CM)),
        {"family": "Blabeloo  MagdySoliman", "size_pt": 16, "fallback_families": ["Tajawal"]},
        resolver=resolver,
    )
    assert with_fallback["substitute"] == "Blabeloo  MagdySoliman"
    assert with_fallback["font_used"]["family"] == "Tajawal" and "البديل" in with_fallback["ar"]


def test_negative_control_missing_glyphs(tmp_path, resolver):
    from fontTools.subset import Options, Subsetter
    from fontTools.ttLib import TTFont

    font = TTFont(str(FONTS / "tajawal" / "Tajawal-Regular.ttf"))
    subsetter = Subsetter(Options())
    subsetter.populate(text="Eyad abc")
    subsetter.subset(font)
    latin = tmp_path / "LatinOnly.ttf"
    font.save(str(latin))
    only = FontResolver(directories=[], files=[latin])
    for method in ("harfbuzz", "pillow_raqm", "approximate"):
        result = fit("إياد", (int(6 * CM), int(3 * CM)), tajawal(16), resolver=only, method=method)
        assert result["verdict"] == "missing_glyphs" and result["notdef"] == 4, method


# --------------------------------------------------------------------------- backends
def test_harfbuzz_and_raqm_agree(resolver):
    box = (int(6 * CM), int(3 * CM))
    rules = {"allow_shrink": True, "min_size_pt": 8}
    hb = fit(BODY, box, tajawal(24), rules, resolver=resolver, method="harfbuzz")
    raqm = fit(BODY, box, tajawal(24), rules, resolver=resolver, method="pillow_raqm")
    assert (hb["size_pt"], hb["lines"]) == (raqm["size_pt"], raqm["lines"])
    assert raqm["method"] == "pillow_raqm" and not raqm["approximate"]


def test_falls_back_when_uharfbuzz_is_missing(resolver, monkeypatch):
    monkeypatch.setitem(sys.modules, "uharfbuzz", None)  # import now raises ImportError
    face = resolver.resolve("Tajawal")
    assert measurer_for(face).method == "pillow_raqm"
    result = fit(BODY, (int(6 * CM), int(3 * CM)), tajawal(16), resolver=resolver)
    assert result["method"] == "pillow_raqm"


def test_approximate_fallback_is_flagged(resolver, monkeypatch):
    monkeypatch.setitem(sys.modules, "uharfbuzz", None)
    monkeypatch.setattr(
        tf_measure.PillowRaqmMeasurer, "__init__", lambda self, face: (_ for _ in ()).throw(ImportError("no raqm"))
    )
    result = fit(BODY, (int(6 * CM), int(3 * CM)), tajawal(16), resolver=resolver)
    assert result["method"] == "approximate" and result["approximate"] and result["measured_with"] == "fontTools-hmtx"
    assert "تقريبي" in result["ar"]
    monkeypatch.setitem(sys.modules, "fontTools.ttLib", None)
    basic = fit(BODY, (int(6 * CM), int(3 * CM)), tajawal(16), resolver=resolver)
    assert basic["method"] == "approximate" and basic["measured_with"] == "Pillow-basic"


# --------------------------------------------------------------------------- PPTX boxes
def _deck(tmp_path: Path, auto_size: str = "none") -> tuple[Path, int]:
    pptx = pytest.importorskip("pptx")
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    prs = pptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(1 * CM), Emu(1 * CM), Emu(int(13.8 * CM)), Emu(int(4.32 * CM)))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = {
        "none": MSO_AUTO_SIZE.NONE,
        "shape": MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT,
        "norm": MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE,
    }[auto_size]
    para = frame.paragraphs[0]
    para.line_spacing = 1.15
    run = para.add_run()
    run.text = DEDICATION
    run.font.size = Pt(24)
    run.font.bold = False
    rpr = run._r.get_or_add_rPr()
    from pptx.oxml.ns import qn

    cs = rpr.makeelement(qn("a:cs"), {"typeface": "Tajawal"})
    latin = rpr.makeelement(qn("a:latin"), {"typeface": "Tajawal"})
    rpr.append(latin)
    rpr.append(cs)
    second = frame.add_paragraph()
    second.text = "مع حبنا"
    plain = slide.shapes.add_textbox(Emu(1 * CM), Emu(8 * CM), Emu(6 * CM), Emu(2 * CM))
    plain.text_frame.text = "إياد"
    out = tmp_path / f"editor_{auto_size}.pptx"
    prs.save(out)
    return out, box.shape_id


def test_paragraph_box_reads_size_insets_and_font(tmp_path):
    deck, shape_id = _deck(tmp_path)
    info = paragraph_box(deck, f"t-2-{shape_id}-0")
    assert info["box_emu"] == (int(13.8 * CM), int(4.32 * CM))
    assert info["inset_emu"] == (91440, 45720, 91440, 45720)
    assert info["font"]["family"] == "Tajawal" and info["font"]["size_pt"] == 24
    assert info["font"]["line_spacing"] == pytest.approx(1.15)
    assert info["text"] == DEDICATION and info["shape_paragraphs"] == [DEDICATION, "مع حبنا"]
    assert info["char_limit"] == char_limit(DEDICATION)


def test_inherited_theme_font_resolves(tmp_path):
    deck, shape_id = _deck(tmp_path)
    info = paragraph_box(deck, f"t-2-{shape_id + 1}-0")
    assert info["font"]["family"] and not info["font"]["family"].startswith("+")
    assert info["font"]["size_pt"] == 18.0  # presentation default when nothing is declared


def test_fit_paragraph_edit_and_refusal(tmp_path, resolver):
    deck, shape_id = _deck(tmp_path)
    ok = fit_paragraph_edit(deck, f"t-2-{shape_id}-0", "بكل الحب تهدى لـإياد", resolver=resolver)
    assert ok["fits"] and ok["paragraph"]["slide"] == 2
    # Within the editor's character cap (raised via original_text) but physically too long for the box:
    bad = fit_paragraph_edit(
        deck, f"t-2-{shape_id}-0", LONG_DEDICATION[:200], resolver=resolver, original_text="x" * 1000
    )
    assert bad["verdict"] == "does_not_fit" and bad["reason"] == "height" and bad["ar"] == TOO_LONG_AR
    # Over the editor's cap for this paragraph (2 x 26 + 40 = 92 characters): the cap is named first.
    capped = fit_paragraph_edit(deck, f"t-2-{shape_id}-0", LONG_DEDICATION[:120], resolver=resolver)
    assert capped["reason"] == "too_many_chars" and capped["char_limit"] == 92
    assert capped["ar"] == "النص ده 120 حرف والمسموح 92 بس — قصّره شوية"


def test_autofit_modes_become_rules(tmp_path, resolver):
    grow, shape_id = _deck(tmp_path, "shape")
    info = paragraph_box(grow, f"t-2-{shape_id}-0")
    assert info["autofit"] == "shape" and info["rules"]["grow_height"]
    long = LONG_DEDICATION[:200]
    assert fit_paragraph_edit(grow, f"t-2-{shape_id}-0", long, resolver=resolver, original_text="x" * 1000)["fits"]
    shrink, shape_id = _deck(tmp_path, "norm")
    info = paragraph_box(shrink, f"t-2-{shape_id}-0")
    assert info["autofit"] == "norm" and info["rules"]["allow_shrink"]
    result = fit_paragraph_edit(
        shrink, f"t-2-{shape_id}-0", ("بكل الحب تهدى لـإياد " * 4).strip(), resolver=resolver, rules={"min_size_pt": 12}
    )
    assert result["verdict"] == "will_shrink" and 12 <= result["size_pt"] < 24


def test_negative_control_unknown_paragraph_ids(tmp_path):
    deck, shape_id = _deck(tmp_path)
    with pytest.raises(ParagraphNotFound):
        paragraph_box(deck, f"t-9-{shape_id}-0")
    with pytest.raises(ParagraphNotFound):
        paragraph_box(deck, f"t-2-{shape_id}-7")
    with pytest.raises(ValueError):
        paragraph_box(deck, "p-2-3-0")


def test_rendered_story_deck_body_paragraph(render_deck, resolver):
    """The Phase-1 synthetic A5 deck: the body paragraph is read as Tajawal 16 pt, 1.15 spacing, and fits."""
    pptx = pytest.importorskip("pptx")
    deck, _spec = render_deck("ar_boy", "A5")
    prs = pptx.Presentation(deck)
    body = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip())
    info = paragraph_box(deck, f"t-1-{body.shape_id}-0")
    assert info["font"]["family"] == "Tajawal" and info["font"]["size_pt"] == 16
    assert info["font"]["line_spacing"] == pytest.approx(1.15)
    assert fit_paragraph_edit(deck, f"t-1-{body.shape_id}-0", info["text"], resolver=resolver)["fits"]
    longer = fit_paragraph_edit(deck, f"t-1-{body.shape_id}-0", (info["text"] + " ") * 8, resolver=resolver)
    assert longer["verdict"] == "does_not_fit"
