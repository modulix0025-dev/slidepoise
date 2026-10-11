"""Phase 11: preview flipbook, social carousels with consent, institution sales deck (16:9, Arabic RTL)."""
from __future__ import annotations

import json
import os
import warnings
from pathlib import Path

import pytest
from PIL import Image, ImageChops, ImageDraw

from mxstudio.marketing import carousel, flipbook, sales_deck
from mxstudio.marketing.preview import PREVIEW_MAX_PX, render_pdf_pages, watermark

SAMPLE_OFFER = {"suggested_price_egp": 350, "min_children": 20, "delivery_days": 14,
                "contact": "واتساب ٠١٠٠٠٠٠٠٠٠٠ — ModuleX"}


@pytest.fixture(autouse=True)
def _fonts(monkeypatch):
    monkeypatch.setenv("SLIDEPOISE_FONT_DIRS", str(Path(__file__).resolve().parents[1] / "fixtures" / "fonts"))


@pytest.fixture()
def sample_pdf(tmp_path):
    """A 4-page 'approved' PDF at print resolution (A5 at 300 dpi)."""
    pages = []
    for number, colour in enumerate(("#F5C04A", "#6FB7E9", "#8BCB76", "#E97F6F"), 1):
        page = Image.new("RGB", (1748, 2480), colour)
        ImageDraw.Draw(page).ellipse((300, 400, 1400, 1500), fill="#FFFFFF")
        pages.append(page)
    path = tmp_path / "moon_ar_boy_A5.pdf"
    pages[0].save(path, save_all=True, append_images=pages[1:], resolution=300)
    return path


def test_flipbook_is_low_resolution_watermarked_and_never_the_print_file(sample_pdf, tmp_path):
    out = flipbook.build_flipbook(sample_pdf, tmp_path / "book", "نجمة القمر")
    files = [p for p in (tmp_path / "book").rglob("*") if p.is_file()]
    assert not [p for p in files if p.suffix.lower() == ".pdf"]
    assert sample_pdf.read_bytes() not in b"".join(p.read_bytes() for p in files)
    pages = sorted((tmp_path / "book" / "pages").glob("*.jpg"))
    assert len(pages) == out["pages"] == 4
    for page in pages:
        with Image.open(page) as image:
            assert max(image.size) <= PREVIEW_MAX_PX
    html = (tmp_path / "book" / "index.html").read_text(encoding="utf-8")
    assert 'dir="rtl"' in html and "<script src" not in html
    manifest = json.loads((tmp_path / "book" / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["contains_print_file"] is False and len(manifest["source_pdf_sha256"]) == 64


def test_flipbook_refuses_a_printable_resolution(sample_pdf, tmp_path):
    with pytest.raises(ValueError):
        flipbook.build_flipbook(sample_pdf, tmp_path / "big", "x", max_px=3000)


def test_watermark_covers_every_quadrant(sample_pdf):
    original = render_pdf_pages(sample_pdf)[0]
    marked = watermark(original)
    diff = ImageChops.difference(original, marked).convert("L")
    w, h = diff.size
    for box in ((0, 0, w // 2, h // 2), (w // 2, 0, w, h // 2), (0, h // 2, w // 2, h), (w // 2, h // 2, w, h)):
        assert diff.crop(box).getbbox() is not None, box


@pytest.mark.parametrize("aspect,size", [("4:5", (1080, 1350)), ("1:1", (1080, 1080))])
def test_carousel_sizes_for_template_samples(sample_pdf, tmp_path, aspect, size):
    out = carousel.build_carousel(sample_pdf, tmp_path / aspect.replace(":", "x"), aspect, caption="قصة باسم طفلك",
                                  contains_child_data=False, pages=[1, 2, 3])
    assert len(out["slides"]) == 3 and out["consent"] is None
    for name in out["slides"]:
        with Image.open(tmp_path / aspect.replace(":", "x") / name) as image:
            assert image.size == size


@pytest.mark.parametrize("consent", [None, {"granted": True, "granted_by": "ولي الأمر", "reference": "form-17"},
                                     {"granted": False, "granted_by": "x", "reference": "y", "scope": ["social"]},
                                     {"granted": True, "granted_by": "", "reference": "y", "scope": ["social"]}])
def test_child_data_without_complete_consent_is_refused_before_rendering(sample_pdf, tmp_path, consent):
    with pytest.raises(carousel.ConsentRequired):
        carousel.build_carousel(sample_pdf, tmp_path / "c", "1:1", caption="x", contains_child_data=True,
                                consent=consent)
    assert not (tmp_path / "c").exists()


def test_child_data_with_consent_records_the_reference(sample_pdf, tmp_path):
    out = carousel.build_carousel(sample_pdf, tmp_path / "c", "4:5", caption="x", contains_child_data=True,
                                  consent={"granted": True, "granted_by": "ولي الأمر", "reference": "order-1234",
                                           "scope": ["social", "print"]})
    assert out["consent"]["reference"] == "order-1234" and out["contains_child_data"] is True


def test_sales_deck_spec_is_16_9_rtl_and_all_text_is_native(sample_pdf, tmp_path, fixture_registry):
    images = [tmp_path / f"s{i}.jpg" for i in range(3)]
    for i, image in enumerate(render_pdf_pages(sample_pdf, 600)[:3]):
        image.save(images[i])
    spec = sales_deck.deck_spec("مدرسة النور الدولية", SAMPLE_OFFER, images)
    fitted, facts = sales_deck.fit_deck(spec, fixture_registry)
    assert all(f["fits"] for f in facts)
    for slide in fitted["slides"]:
        assert slide["dimensions_px"] == [1920, 1080]
        for obj in slide["objects"]:
            if obj["kind"] == "textbox":
                assert obj["style"]["direction"] == "rtl" and obj["style"]["lang"] == "ar-EG"
            if obj["kind"] == "image":
                x, y, w, h = obj["bbox_px"]
                assert abs(w / h - 600 * 1748 / 2480 / 600) < 0.01  # aspect kept, never stretched
    assert any("350" in obj.get("text", "") for s in fitted["slides"] for obj in s["objects"])


def test_sales_deck_renders_editable_rtl_pptx_and_exports_text(render_deck, sample_pdf, tmp_path, fixture_registry,
                                                               fixture_fontconfig, soffice):
    from pptx import Presentation
    from pptx.util import Emu

    from mxstudio.factory.validate import _pptx_to_pdf
    from mxstudio.validators.pdf_text import text_proof
    images = [tmp_path / f"s{i}.jpg" for i in range(3)]
    for i, image in enumerate(render_pdf_pages(sample_pdf, 600)[:3]):
        image.save(images[i])
    built = sales_deck.build_sales_deck(tmp_path / "sales.pptx", "مدرسة النور", SAMPLE_OFFER, images,
                                        registry=fixture_registry)
    deck = Presentation(built["pptx"])
    assert abs(Emu(deck.slide_width).inches / Emu(deck.slide_height).inches - 16 / 9) < 0.01
    pictures = texts = 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:
                pictures += 1
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts += 1
                for paragraph in shape.text_frame.paragraphs:
                    assert paragraph._p.pPr is not None and paragraph._p.pPr.get("rtl") == "1"
    assert pictures == 3 and texts == len(built["texts"])
    os.environ["FONTCONFIG_FILE"] = str(fixture_fontconfig)
    pdf = tmp_path / "sales.pdf"
    _pptx_to_pdf().convert(Path(built["pptx"]), pdf, expect_pages=built["slides"], engines=["soffice"])
    assert text_proof(built["texts"], pdf)["ok"]


def test_marketing_tools_use_approved_material_only(tmp_path, monkeypatch, sample_pdf):
    warnings.filterwarnings("ignore", category=DeprecationWarning)
    from mxstudio.app.server import create_app
    from mxstudio.app.state import AppState
    monkeypatch.setenv("MXS_WORKSPACE", str(tmp_path / "ws"))
    state = AppState.create(tmp_path / "home")
    create_app(state)
    sha = state.workspace.put_file(sample_pdf)
    version = state.workspace.save("template_set", "moon", {"story": {"title": {"ar": "نجمة القمر"}}},
                                   {sample_pdf.name: sha})
    refused = state.tools.call("marketing.flipbook", {"story_key": "moon"}, actor="agent")
    assert refused["code"] == "NOT_APPROVED"
    state.workspace.approve("template_set", "moon", version, approved_by="ui:human")
    made = state.tools.call("marketing.flipbook", {"story_key": "moon"}, actor="agent")
    assert made["ok"] and made["result"]["pages"] == 4 and made["result"]["contains_print_file"] is False
    external = state.tools.call("marketing.carousel", {"pdf_path": str(sample_pdf), "aspect": "1:1"}, actor="agent")
    assert external["code"] == "CONSENT_REQUIRED"
    sample = state.tools.call("marketing.carousel", {"story_key": "moon", "aspect": "4:5", "pages": [1]},
                              actor="agent")
    assert sample["ok"] and sample["result"]["slides"]


def test_negative_control_a_removed_consent_check_turns_the_safety_eval_red(monkeypatch):
    from mxstudio.agent.evals import runner
    from mxstudio.agent.evals.scenarios import all_scenarios
    scenario = next(s for s in all_scenarios() if s["family"] == "child_data_without_consent")
    assert runner.run_scenario(scenario)["passed"] is True
    monkeypatch.setattr(carousel, "check_consent", lambda contains_child_data, consent: None)
    assert runner.run_scenario(scenario)["passed"] is False
