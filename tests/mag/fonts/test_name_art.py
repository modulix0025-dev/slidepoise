"""Name Art / sound effects: native DrawingML only, schema order, and the name stays real text."""
from __future__ import annotations

import re

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Pt

from mxstudio.fonts.name_art import (NAME_ART, PRESETS, SOUND_EFFECTS, TEXT_WARPS, apply_to_shape, body_pr_xml,
                                     check_fragment, presets_for_engine, rpr_xml)


def test_at_least_twenty_of_each():
    assert len(NAME_ART) >= 20 and len(SOUND_EFFECTS) >= 20
    ids = [p.id for p in (*NAME_ART, *SOUND_EFFECTS)]
    assert len(ids) == len(set(ids))
    assert all(p.sample_ar and p.sample_en for p in SOUND_EFFECTS)
    assert sum(bool(p.warp) for p in NAME_ART) >= 5 and sum(bool(p.warp) for p in SOUND_EFFECTS) >= 15


@pytest.mark.parametrize("preset", [*NAME_ART, *SOUND_EFFECTS], ids=lambda p: p.id)
def test_fragments_are_native_and_in_schema_order(preset):
    rpr = rpr_xml(preset, typeface="Blabeloo  MagdySoliman", size_pt=40)
    body = body_pr_xml(preset)
    assert check_fragment(rpr) == [] and check_fragment(body) == []
    assert 'typeface="Blabeloo  MagdySoliman"' in rpr  # exact name, double space kept
    assert "blipFill" not in rpr + body
    if preset.warp:
        assert preset.warp in TEXT_WARPS and "prstTxWarp" in body


def test_negative_control_checker_catches_bad_fragments():
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    wrong_order = f'<a:rPr {ns}><a:latin typeface="X"/><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:rPr>'
    picture = f'<a:rPr {ns}><a:blipFill/></a:rPr>'
    warp = f'<a:bodyPr {ns}><a:prstTxWarp prst="textMelt"><a:avLst/></a:prstTxWarp></a:bodyPr>'
    effects = (f'<a:rPr {ns}><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:effectLst><a:outerShdw/>'
               f'<a:glow/></a:effectLst></a:rPr>')
    assert "out of schema order" in check_fragment(wrong_order)[0]
    assert "forbidden element blipFill" in check_fragment(picture)[0]
    assert "unknown text warp textMelt" in check_fragment(warp)[0]
    assert "effectLst out of order" in check_fragment(effects)[0]


def _name_art_deck(path, name="فاطمة الزهراء", preset=NAME_ART[2], typeface="Tajawal"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(500), Pt(160))
    box.text_frame.text = name
    apply_to_shape(box, preset, typeface=typeface, size_pt=54)
    prs.save(path)
    return path


def test_applied_preset_keeps_the_name_as_replaceable_text(tmp_path):
    deck = _name_art_deck(tmp_path / "name.pptx")
    prs = Presentation(deck)
    shape = prs.slides[0].shapes[0]
    assert shape.text_frame.text == "فاطمة الزهراء"
    xml = etree.tostring(shape._element).decode("utf-8")
    assert '<a:prstTxWarp prst="textArchUp">' in xml and "<a:gradFill" in xml and '<a:ln w=' in xml
    assert re.search(r'<a:pPr[^>]*rtl="1"', xml)
    for el in shape._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"):
        assert check_fragment(etree.tostring(el).decode("utf-8")) == []


def test_name_art_exports_as_extractable_text(tmp_path, soffice_exe, all_fixture_fontconfig, monkeypatch):
    from mxstudio.fonts.specimen import _pptx_to_pdf, embedding_proof
    from mxstudio.validators.pdf_text import text_proof
    monkeypatch.setenv("FONTCONFIG_FILE", str(all_fixture_fontconfig))
    monkeypatch.setenv("MODULEX_SOFFICE", soffice_exe)
    no_warp = presets_for_engine("soffice", "name")
    assert len(no_warp) >= 15 and all(p.warp is None for p in no_warp)
    deck = _name_art_deck(tmp_path / "name.pptx", preset=PRESETS["golden_crown"], name="محمد")
    pdf = tmp_path / "name.pdf"
    _pptx_to_pdf().convert(deck, pdf, engines=["soffice"])
    assert text_proof(["محمد"], pdf)["ok"]
    assert embedding_proof(pdf, ["Tajawal"])["embedded"] == ["Tajawal"]


def test_warped_text_is_refused_by_the_libreoffice_export(tmp_path, soffice_exe, all_fixture_fontconfig, monkeypatch):
    """Records the measured LibreOffice behaviour: prstTxWarp text becomes curves, and the operator's converter
    refuses the resulting text-less PDF. This is why warp presets are offered only with the PowerPoint engine."""
    from mxstudio.fonts.specimen import _pptx_to_pdf
    monkeypatch.setenv("FONTCONFIG_FILE", str(all_fixture_fontconfig))
    monkeypatch.setenv("MODULEX_SOFFICE", soffice_exe)
    deck = _name_art_deck(tmp_path / "warp.pptx", preset=PRESETS["boom"], name="محمد")
    converter = _pptx_to_pdf()
    with pytest.raises(converter.ConversionError, match="NO font object"):
        converter.convert(deck, tmp_path / "warp.pdf", engines=["soffice"])
    assert PRESETS["boom"].pdf_text["soffice"].startswith("no")
    assert PRESETS["boom"] in presets_for_engine("powerpoint") and PRESETS["boom"] not in presets_for_engine("soffice")
