"""Deck font validator: every exported typeface installed AND certified, refused by its exact name."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Pt

from mxstudio.fonts.deck_validator import certified_install_names, referenced_typefaces, validate_deck

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
DOUBLE = "Blabeloo  MagdySoliman"


def _deck(tmp_path: Path, runs: list[tuple[str, str]], *, notes_font: str | None = None,
          theme=("Tajawal", "Tajawal", "Amiri")) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, Pt(300), Pt(100))
    for text, face in runs:
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        rpr = run._r.get_or_add_rPr()
        for tag in ("latin", "cs"):
            node = rpr.makeelement(f"{{{A}}}{tag}", {"typeface": face})
            rpr.append(node)
    if notes_font:
        notes = slide.notes_slide.notes_text_frame
        notes.text = "note"
        rpr = notes.paragraphs[0].runs[0]._r.get_or_add_rPr()
        rpr.append(rpr.makeelement(f"{{{A}}}latin", {"typeface": notes_font}))
    raw = tmp_path / "raw.pptx"
    prs.save(raw)
    out = tmp_path / "deck.pptx"
    major, minor, arab = theme
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r'(<a:majorFont><a:latin typeface=")[^"]*', rf"\g<1>{major}", xml)
                xml = re.sub(r'(<a:minorFont><a:latin typeface=")[^"]*', rf"\g<1>{minor}", xml)
                if arab is not None:
                    xml = re.sub(r'<a:font script="Arab" typeface="[^"]*"/>', f'<a:font script="Arab" typeface="{arab}"/>',
                                 xml)
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    return out


def test_theme_tokens_resolve_through_the_master_theme(tmp_path):
    deck = _deck(tmp_path, [("مرحبا", "Tajawal")])
    exported, _hidden, unresolved = referenced_typefaces(deck)
    resolved = {(r.via_theme, r.typeface) for r in exported if r.via_theme}
    assert ("+mn-lt", "Tajawal") in resolved and ("+mj-lt", "Tajawal") in resolved
    assert ("+mn-cs", "Amiri") in resolved  # empty theme cs → <a:font script="Arab">
    assert {u["token"] for u in unresolved} == {"+mj-ea", "+mn-ea"}


def test_installed_and_certified_deck_passes(tmp_path, fixture_registry):
    deck = _deck(tmp_path, [("مرحبا ", "Tajawal"), ("يا ", "Amiri")])
    report = validate_deck(deck, installed=fixture_registry, certified={"Tajawal", "Amiri"})
    assert report.ok, report.refusals
    assert {t["name"] for t in report.typefaces} == {"Tajawal", "Amiri"}


def test_negative_control_uninstalled_font_refused_by_exact_name(tmp_path, fixture_registry):
    deck = _deck(tmp_path, [("لـيوسف", DOUBLE), ("مرحبا", "Tajawal")])
    report = validate_deck(deck, installed=fixture_registry, certified={"Tajawal", "Amiri", DOUBLE})
    assert not report.ok
    assert [(r["finding"], r["typeface"]) for r in report.refusals] == [("deck_font_not_installed", DOUBLE)]
    assert f"«{DOUBLE}»" in report.refusals[0]["ar"]


def test_double_space_name_is_not_the_single_space_name(tmp_path):
    deck = _deck(tmp_path, [("لـيوسف", DOUBLE)])
    base = {"Tajawal", "Amiri"}
    single = validate_deck(deck, installed=base | {"Blabeloo MagdySoliman"}, certified=base | {"Blabeloo MagdySoliman"})
    assert [r["typeface"] for r in single.refusals] == [DOUBLE]
    uncertified = validate_deck(deck, installed=base | {DOUBLE}, certified=base)
    assert [(r["finding"], r["typeface"]) for r in uncertified.refusals] == [("deck_font_not_certified", DOUBLE)]
    ok = validate_deck(deck, installed=base | {DOUBLE}, certified=base | {DOUBLE})
    assert ok.ok and DOUBLE in {t["name"] for t in ok.typefaces}


def test_negative_control_default_theme_fonts_are_refused(tmp_path, fixture_registry):
    deck = _deck(tmp_path, [("hi", "Tajawal")], theme=("Calibri", "Calibri", None))  # python-pptx default theme
    report = validate_deck(deck, installed=fixture_registry, certified={"Tajawal", "Amiri"})
    # +mj-cs → <a:font script="Arab"> of majorFont (Times New Roman), +mn-cs → minorFont's (Arial)
    assert {r["typeface"] for r in report.refusals} == {"Calibri", "Times New Roman", "Arial"}


def test_notes_are_listed_but_not_exported(tmp_path, fixture_registry):
    deck = _deck(tmp_path, [("hi", "Tajawal")], notes_font="Comic Sans MS")
    report = validate_deck(deck, installed=fixture_registry, certified={"Tajawal", "Amiri"})
    assert report.ok
    assert "Comic Sans MS" in {n["name"] for n in report.not_exported}


def test_certified_install_names_use_renamed_builds(committed):
    names = certified_install_names()
    assert {"MX Kids F033 Readex Pro", "MX Kids F056", "Tajawal", "Baloo Bhaijaan 2"} <= names
    assert "Readex Pro" not in names and "Nabla" not in names
    assert DOUBLE in certified_install_names(extra=[DOUBLE])
