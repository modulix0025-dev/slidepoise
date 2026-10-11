"""Synthetic known-answer tests for the Phase 0 anatomy tool (no operator data needed)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import mag_template_anatomy as anatomy  # noqa: E402

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _fixture(tmp: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(14.8), Cm(21.0)
    picture_layout = prs.slide_layouts[8]  # "Picture with Caption": has a pic placeholder
    s1 = prs.slides.add_slide(picture_layout)  # empty pic placeholder stays an <p:sp>
    s2 = prs.slides.add_slide(picture_layout)
    art = tmp / "art.png"
    Image.new("RGB", (437, 620), "white").save(art)  # 437 px over 14.8 cm = 75 dpi
    for ph in s2.placeholders:
        if ph.placeholder_format.type == 18:  # PICTURE -> filled, must not count
            ph.insert_picture(str(art))
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_picture(str(art), 0, 0, prs.slide_width, prs.slide_height)
    box = s3.shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(3))
    para = box.text_frame.paragraphs[0]
    for text in ("مرحبا يا ", "يو", "سف"):  # base_name split over two runs
        run = para.add_run()
        run.text = text
    para._p.get_or_add_pPr().set("rtl", "1")
    para2 = box.text_frame.add_paragraph()
    run = para2.add_run()
    run.text = "يوسف بطل القصة"
    rpr = run._r.get_or_add_rPr()
    etree.SubElement(rpr, f"{{{A}}}cs").set("typeface", "Blabeloo  MagdySoliman")
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    s4._element.set("show", "0")
    out = tmp / "a5 test.pptx"
    prs.save(out)
    return out


def test_known_answers(tmp_path):
    path = _fixture(tmp_path)
    record = anatomy.measure(path, "يوسف", None)
    assert record["slide_size_cm"] == [14.8, 21.0]
    assert record["slide_count"] == 4
    assert record["hidden_slides"] == [4]
    assert record["picture_slots_local"] == {1: 1, 2: 0, 3: 0, 4: 0}
    assert record["base_name_total"] == 2
    assert record["base_name_split_across_runs"] == 1
    assert "Blabeloo  MagdySoliman" in record["typefaces"]  # double space preserved
    image = record["slides"][2]["images"][0]
    assert image["effective_dpi"] == 75.0
    assert all(image["touches_trim"].values())
    para = record["slides"][2]["text_shapes"][0]["paragraphs"][0]
    assert para["rtl"] == "1"


def test_cli_writes_json(tmp_path):
    _fixture(tmp_path)
    out = tmp_path / "anatomy.json"
    assert anatomy.main(["--templates", str(tmp_path), "--output", str(out), "--base-name", "يوسف"]) == 0
    report = json.loads(out.read_text(encoding="utf-8"))
    assert report["count"] == 1
    assert report["picture_slot_counter"] == "local_reimplementation_unverified"
    assert report["templates"][0]["base_name_used"] == "يوسف"
